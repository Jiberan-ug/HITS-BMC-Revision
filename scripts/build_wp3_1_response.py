#!/usr/bin/env python3
"""Create a verbatim point-by-point response and comment coverage audit."""

from __future__ import annotations

import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


REPO = Path(__file__).resolve().parents[1]
OUT = REPO / "HITS_BMC_MINOR_REVISION_WP3_1_SUBMISSION_READY"
EDITOR = Path.home() / ".codex_fast/attachments/d1a85934-b384-4434-9f1e-049d7337c5a7/已粘贴的文本.txt"
REVIEWERS = Path.home() / ".codex_fast/attachments/ff8bc8db-e06e-4e43-9dba-62965973e4ec/已粘贴的文本.txt"
MANUSCRIPT = OUT / "01_MANUSCRIPT/HITS_BMC_Revised_Manuscript_CLEAN.docx"
PDF = OUT / "01_MANUSCRIPT/HITS_BMC_Revised_Manuscript_CLEAN.pdf"
RESPONSE = OUT / "02_RESPONSE/HITS_BMC_Point_by_Point_Response_FINAL.docx"


def numbered_comments(text: str) -> list[str]:
    found = re.findall(r"(?ms)^\s*(\d+)[.。]\s*(.*?)(?=^\s*\d+[.。]\s*|\Z)", text)
    return [body.strip().strip('"') for _, body in found]


def read_comments():
    editor_text = EDITOR.read_text(encoding="utf-8")
    editor_area = editor_text.split("Editor Comments", 1)[1].split("\n\nWith above comments kindly also address:", 1)[0]
    editor_comments = numbered_comments(editor_area)
    reviewer_text = REVIEWERS.read_text(encoding="utf-8")
    r1_text = reviewer_text.split("Reviewer 1", 1)[1].split("Reviewer 2", 1)[0]
    major_area = r1_text.split("Major comments", 1)[1].split("Minor comments", 1)[0]
    minor_area = r1_text.split("Minor comments", 1)[1]
    r1_major = numbered_comments(major_area)
    r1_minor = numbered_comments(minor_area)
    r2_text = reviewer_text.split("Reviewer 2", 1)[1].strip()
    if (len(editor_comments), len(r1_major), len(r1_minor)) != (10, 7, 3):
        raise RuntimeError(f"Unexpected comment counts: editor={len(editor_comments)}, R1 major={len(r1_major)}, R1 minor={len(r1_minor)}")
    if not r2_text:
        raise RuntimeError("Reviewer 2 comment is missing")
    ethics = editor_text.split("With above comments kindly also address:", 1)[1].split("-Youfu He", 1)[0].strip().strip('"')
    return editor_comments, ethics, r1_major, r1_minor, r2_text


REPLIES = {
    ("Editor", "1"): (
        "We clarified the record-selection sequence: a non-empty discharge diagnosis was prioritized, followed by the number of available core absolute CBC counts, a parsable CBC/WBC timestamp, and the overall non-missing-field count, with source-row order as the final tie-break. Admission, angiography, diagnosis, and treatment timestamps could not be linked consistently to the selected measurements. We therefore do not describe them as the first admission CBC or as pre-angiography or pre-diagnosis measurements.",
        "The CBC selection rule and timing limitations are stated in Methods; the unresolved source-date limitation is disclosed in the Discussion.",
        "the Core-7 predictors were absolute neutrophil",
        "Predictors and laboratory measurements"),
    ("Editor", "2"): (
        "AMI phenotype was classified from discharge-diagnosis text. Independent clinical adjudication was not available. Cardiac injury markers were supportive only because assay reference limits, serial changes, and measurement timing were incomplete; they were not used to redefine the full cohort. We have clarified these limits and distinguished the strict discharge-text sensitivity definition from clinical adjudication.",
        "The phenotype source, operational definitions, supportive role of biomarkers, and lack of independent adjudication are specified in Methods and reiterated in the Discussion.",
        "AMI phenotype was classified from discharge-diagnosis text",
        "Outcome definition"),
    ("Editor", "3"): (
        "The source export contained 2,548 records representing 2,279 unique patients. Of these, 2,010 had one source record and 269 had two; thus 269 surplus records were removed when retaining one patient-level record per person. The source does not establish that these were distinct admissions. The retained record followed the stated completeness-based priority rule. Figure 1 now distinguishes source records, unique patients, phenotype groups, and the final analysis cohort.",
        "Methods and Figure 1 report the record-to-patient reconciliation and selection rule; the figure does not label the records as repeat admissions.",
        "The source export contained 2,548 records",
        "Eligibility and study population"),
    ("Editor", "4"): (
        "We clarified that stratified five-fold outer cross-validation was repeated 10 times, with five-fold tuning within each outer training set. Each patient’s 10 held-out probabilities were averaged for evaluation. Paired models used identical outer partitions where applicable. Imputation, transformations, scaling, and tuning were confined to training data. The 1,000 bootstrap resamples used fixed patient-level mean predictions and outcomes without refitting; accordingly, the intervals do not capture full model-development uncertainty.",
        "The aggregation, shared partitions, training-only processing, and fixed-prediction bootstrap are described in Methods.",
        "Primary internal validation used stratified outer five-fold",
        "Internal validation and uncertainty"),
    ("Editor", "5"): (
        "The date-axis split point was 1 January 2015; the development group included measurements dated on or before 2 July 2014, the buffer ran from 3 July 2014 through 2 July 2015, and the later group began on 3 July 2015. The earlier group was used for development, the later group for fixed evaluation, and the buffer was excluded. The groups were development N=1,001 (AMI=196), buffer N=271 (AMI=79), and later N=548 (AMI=178). In the later group, the Core-7-minus-PIV ΔAUC was 0.002 (95% CI -0.021 to 0.025), which includes zero. We now describe this as an exploratory deidentified date-axis sensitivity analysis, not formal temporal validation. A source audit also identified a material unresolved discrepancy: selected CBC timestamps span 2009-2026, with 1,802/1,820 before 2019, whereas the manuscript states a 2020-2026 study period. The documented bounded shift and linkage to the analytic input do not resolve this; we have not interpreted the date axis as actual calendar-time transportability.",
        "The split, buffer, counts, estimate, and limitation are reported in Methods and Results. The source-date discrepancy remains flagged for author/institutional reconciliation before submission.",
        "In the later group, Core-7, PIV, and Enhanced AUCs were",
        "Exploratory date-axis and diagnostic sensitivity analyses"),
    ("Editor", "6"): (
        "The high-specificity comparison retained 138 strict-text AMI cases from the 178 later-group AMI cases and included 288 strict CAD/angina controls (N=426). The other 40 later-group AMI cases did not meet the text rule; they are not described as adjudicated misclassifications. This was a discharge-text sensitivity analysis rather than independent clinical adjudication. A separate flowchart is provided in Additional file 1, Figure S2.",
        "The selection and interpretation are clarified in Results; the supplementary flowchart shows the high-specificity comparison.",
        "The high-specificity comparison contained",
        "Exploratory date-axis and diagnostic sensitivity analyses"),
    ("Editor", "7"): (
        "Fibrinogen was missing in 115/1,820 patients (6.32%), including 27/453 AMI patients (5.96%) and 88/1,367 non-AMI CAD patients (6.44%). Core-7 and Enhanced were compared in the same fibrinogen-complete 1,705 patients and on identical validation partitions; fibrinogen was not imputed for this comparison.",
        "Methods reports group-specific missingness; Results reports the same-sample complete-case comparison on identical validation partitions.",
        "Fibrinogen was missing in 115/1,820 patients",
        "Phenotype and missingness sensitivity analyses"),
    ("Editor", "8"): (
        "On the same fibrinogen-complete sample (N=1,705; AMI=426), AUCs were 0.546 for Clinical, 0.717 for Clinical+PIV, 0.730 for Clinical+Core-7, and 0.741 for Clinical+Enhanced. Corresponding Brier scores were 0.187, 0.166, 0.162, and 0.161; the Clinical+Core-7 versus Clinical+PIV paired ΔAUC was 0.0130 (95% CI 0.0004-0.0251). Age and sex were complete; hypertension was missing in 79/1,820 (4.34%) and diabetes in 48/1,820 (2.64%), with both imputed using training-fold medians. Smoking was unavailable. The full paired comparisons are in Additional file 2, Table S8.",
        "The same-sample four-model results, clinical missingness, and paired interval are reported in Methods/Results and Additional file 2, Table S8.",
        "The prespecified clinical baseline model included",
        "Clinical incremental value and fibrinogen enhancement"),
    ("Editor", "9"): (
        "Calibration intercepts, slopes, and plots were calculated from the patient-level arithmetic mean of the 10 held-out predictions. Additional file 2, Table S6 reports the calibration-group sizes, AMI counts, and predicted-probability summaries. The available analysis output provided some alternative date-shift estimates only as point estimates; confidence intervals were unavailable, so we report them as such rather than implying interval estimation.",
        "The prediction source and group-level counts are identified in Results and Additional file 2, Table S6; Figure 5 and its legend identify estimates without intervals.",
        "Calibration intercept and slope were",
        "Internal model performance"),
    ("Editor", "10"): (
        "We revised the abstract, Discussion, and Conclusions to characterize the gains as modest additional information for AMI phenotype discrimination. The Core-7 increment over PIV was attenuated in the exploratory date-axis comparison, and the findings do not establish diagnostic or triage readiness. The existing qualifications are retained and strengthened where needed.",
        "The revised claim appears in the Abstract, Discussion, and Conclusions.",
        "Multidimensional analysis of routine CBC variables provided",
        "Conclusions"),
    ("Editor ethics", "1"): (
        "We added the requested statement that all methods were performed in accordance with relevant guidelines and regulations. The Ethics section identifies approval no. K202602-10, and the consent statement now records the ethics-committee waiver for this retrospective study.",
        "The Ethics approval and Consent to participate declarations are corrected; Consent for publication is stated as not applicable.",
        "The study was approved by the Ethics Committee",
        "Ethics approval"),
    ("Reviewer 1 Major", "1"): (
        "We specified the nested five-fold outer/inner procedure, training-fold-only preprocessing and tuning, and patient-level averaging of the 10 held-out predictions. No independent external cohort was used. The date-axis analysis is now explicitly exploratory and not formal temporal validation; the stated date-source discrepancy and unresolved linkage are disclosed, so external applicability remains uncertain.",
        "Methods and Discussion describe the validation procedure and limitations; Results reports the fixed later-group comparison.",
        "Primary internal validation used stratified outer five-fold",
        "Internal validation and uncertainty"),
    ("Reviewer 1 Major", "2"): (
        "We report the exploratory correlation assessment and selection stability in Additional file 1, Figure S3 and Additional file 2, Table S7. The maximum absolute pairwise Spearman correlation among the Core-7 predictors was 0.469 (neutrophil and monocyte counts). Neutrophils, lymphocytes, monocytes, platelets, and RDW were selected in all resampling iterations; MPV was selected in 99% and hemoglobin in 95%. Elastic-net estimates are interpreted as predictive associations, not isolated biological effects.",
        "The stability summary is reported in Results and Additional file 2, Table S7; predictor correlations are shown in Additional file 1, Figure S3.",
        "In the primary stability analysis",
        "Predictor stability"),
    ("Reviewer 1 Major", "3"): (
        "Age and sex were complete. Hypertension was missing in 79/1,820 patients (4.34%) and diabetes in 48/1,820 (2.64%). Missing clinical values were imputed using medians estimated from each training fold; smoking was unavailable and was not included. Imputation remained inside the validation pipeline.",
        "Methods and Results specify the clinical-variable completeness and fold-specific imputation.",
        "Hypertension was missing in 79/1,820",
        "Model development and preprocessing"),
    ("Reviewer 1 Major", "4"): (
        "We report AUCs and 95% confidence intervals for all prespecified conventional indices in Additional file 2, Table S4. Each index was evaluated as a continuous predictor in the repeated nested validation framework. PIV was the prespecified primary comparator for the paired Core-7 comparison; the other indices are descriptive benchmarks, and no additional pairwise Core-7-versus-index tests were performed.",
        "The Methods now distinguishes the prespecified PIV comparison from descriptive benchmark indices; all benchmark AUCs and confidence intervals are in Additional file 2, Table S4.",
        "Using absolute counts, we calculated NLR",
        "Conventional inflammatory indices"),
    ("Reviewer 1 Major", "5"): (
        "The source export contained 2,279 unique patients; 269 had two source records, but the data do not establish that these were distinct admissions. One record per patient was selected using the stated completeness-based priority rule. Validated index-relative fields for prior CAD, MI, PCI, or CABG were unavailable, and diagnosis mentions could not establish prior-history prevalence. We therefore did not report unsupported prevalences or include these as predictors; this limitation is now explicit.",
        "The patient-level flow and selection rule are shown in Methods/Figure 1, and unavailable prior-history fields are acknowledged in the Discussion.",
        "The source export contained 2,548 records",
        "Eligibility and study population"),
    ("Reviewer 1 Major", "6"): (
        "PIV was modeled as a single continuous predictor after log(1+PIV) transformation and standardization using training-fold parameters, in an unpenalized logistic model. It was evaluated with the same patients and outer-fold partitions as Core-7 where applicable. No post hoc calibration was applied; all reported performance derives from held-out predictions.",
        "The PIV transformation and shared validation framework are specified in Methods; comparative estimates are reported in Tables 2 and S4.",
        "PIV was the prespecified primary comparator",
        "Conventional inflammatory indices"),
    ("Reviewer 1 Major", "7"): (
        "A standardized incremental comparison beyond high-sensitivity troponin was not possible because assay platforms, reference limits, serial change, and timing were not sufficiently harmonized. Troponin and CK-MB served only as supportive phenotype checks, not as a complete adjudication standard. We now state that the model was not evaluated for real-time triage, treatment decisions, or performance before biomarker results; any future role would require independent, time-anchored evaluation.",
        "The troponin-comparator limitation and absence of diagnostic/triage readiness are stated in the Discussion.",
        "No standardized troponin comparator was available",
        "Discussion"),
    ("Reviewer 1 Minor", "1"): (
        "A complete eligible-patient screening log was unavailable, so consecutive enrollment could not be confirmed. We state this explicitly and discuss possible selection related to data completeness and diagnosis availability. The selected CBC timestamp discrepancy with the stated study period also remains unresolved and is a submission hold.",
        "The study-design paragraph and limitations now state that consecutive screening cannot be confirmed and identify selection concerns.",
        "Because a complete eligible-patient screening log was unavailable",
        "Study design and setting"),
    ("Reviewer 1 Minor", "2"): (
        "The selected record was determined by the stated completeness-based rule, not by a verified clinical-time hierarchy. Encounter-level timing was insufficient to establish whether the selected CBC was first on admission or preceded symptom assessment, troponin testing, angiography, diagnosis, or treatment. We therefore avoid those temporal claims.",
        "The CBC measurement paragraph states the selection rule and timing limitations.",
        "Encounter-level linkage and admission",
        "Predictors and laboratory measurements"),
    ("Reviewer 1 Minor", "3"): (
        "The ethics statement now specifies that informed consent was waived by the Ethics Committee because of the retrospective study design; the prior statement that all participants provided consent has been removed.",
        "Declarations, Ethics approval and Consent to participate.",
        "Given the retrospective nature of the study",
        "Consent to participate"),
    ("Reviewer 2", "1"): (
        "We added exploratory decision-curve analysis for the four clinical models using the same fibrinogen-complete sample (N=1,705) and held-out predictions, with treat-all and treat-none references across thresholds of 0.05-0.50. The analysis is supplementary and is not presented as evidence of clinical utility or treatment benefit.",
        "Methods and Results describe the analysis; the curve is in Additional file 1, Supplementary Figure S1.",
        "An exploratory decision-curve analysis used",
        "Reporting and reproducibility"),
}


def compact(text: str) -> str:
    return re.sub(r"[^a-z]+", "", text.lower())


def page_for_paragraph(paragraph_text: str, anchor: str) -> int:
    from pypdf import PdfReader

    needle = compact(anchor)
    compact_para = compact(paragraph_text)
    pos = compact_para.find(needle)
    before = compact_para[max(0, pos - 35):pos] if pos >= 0 else ""
    after = compact_para[pos + len(needle):pos + len(needle) + 70] if pos >= 0 else ""
    for number, page in enumerate(PdfReader(str(PDF)).pages, 1):
        page_text = compact(page.extract_text() or "")
        if needle in page_text and (not after or after[:35] in page_text or before[-25:] in page_text):
            return number
    for number, page in enumerate(PdfReader(str(PDF)).pages, 1):
        if needle in compact(page.extract_text() or ""):
            return number
    raise RuntimeError(f"Could not locate manuscript anchor in rendered PDF: {anchor}")


def paragraph_location(anchor: str, section_name: str) -> str:
    doc = Document(MANUSCRIPT)
    active_section = None
    active_subsection = None
    count = 0
    target_number = None
    target_text = None
    for para in doc.paragraphs:
        if para.style.name == "Heading 1":
            active_section = para.text.strip()
            active_subsection = None
            count = 0
            continue
        if para.style.name.startswith("Heading"):
            active_subsection = para.text.strip()
            count = 0
            continue
        if active_section and para.text.strip():
            count += 1
            expected_heading = (active_subsection or active_section).lower()
            if anchor.lower() in para.text.lower() and (section_name.lower() in expected_heading or expected_heading in section_name.lower()):
                target_number = count
                target_text = para.text
                break
    if target_number is None:
        raise RuntimeError(f"Could not locate manuscript paragraph anchor: {anchor}")
    page = page_for_paragraph(target_text, anchor)
    return f"page {page}, paragraph {target_number} of {section_name}"


def set_run_font(run, size=10.5, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold


def add_label(doc, label: str, text: str, bold_label=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    a = p.add_run(label)
    set_run_font(a, 10.5, bold_label)
    b = p.add_run(text)
    set_run_font(b)


def main():
    editor_comments, ethics_comment, r1_major, r1_minor, r2_comment = read_comments()
    if not MANUSCRIPT.exists() or not PDF.exists():
        raise FileNotFoundError("Render the clean manuscript to PDF before creating the response letter")
    RESPONSE.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size in (("Title", 16), ("Heading 1", 13), ("Heading 2", 11.5)):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.style = doc.styles["Normal"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    r = title.add_run("Response to the Editor and Reviewers")
    set_run_font(r, 16, True)
    for label, value in (
        ("Manuscript title: ", "Multidimensional routine hematologic modeling for acute myocardial infarction phenotype discrimination in coronary artery disease: a single-center retrospective study"),
        ("Journal: ", "BMC Cardiovascular Disorders"),
        ("Submission ID: ", "fa1e4873-011b-4f15-8b0d-29f6df21fabe"),
    ):
        add_label(doc, label, value)
    intro = doc.add_paragraph("Dear Editor and Reviewers,\n\nWe thank you for the careful assessment. We have made targeted revisions and respond point by point below. The selected CBC date lineage remains under author and institutional verification; the date-axis analysis is therefore described as exploratory rather than temporal validation.")
    intro.paragraph_format.space_after = Pt(10)

    coverage = []

    def add_comment_block(group: str, number: str, heading: str, exact_comment: str, status: str):
        doc.add_heading(heading, level=2)
        add_label(doc, "Comment: ", exact_comment)
        key = (group, number)
        reply, change_text, anchor, section_name = REPLIES[key]
        add_label(doc, "Response: ", reply)
        loc = paragraph_location(anchor, section_name)
        if group == "Editor" and number == "6":
            loc += "; Additional file 1, Figure S2"
        elif group == "Editor" and number == "8":
            loc += "; Additional file 2, Table S8"
            loc += "; " + paragraph_location("To assess whether hematologic information added", "Model development and preprocessing")
        elif group == "Editor" and number == "9":
            loc += "; Additional file 2, Table S6"
        elif group == "Editor" and number == "1":
            loc += "; " + paragraph_location("available selected CBC timestamps spanned 2009-2026", "Discussion")
        elif group == "Editor" and number == "5":
            loc += "; " + paragraph_location("To comply with institutional privacy regulations, exact calendar dates", "Exploratory deidentified date-axis sensitivity analysis")
        elif group == "Reviewer 1 Major" and number == "2":
            loc += "; Additional file 1, Figure S3; Additional file 2, Table S7"
        elif group == "Reviewer 1 Major" and number == "3":
            loc += "; " + paragraph_location("Clinical covariate missingness in the primary cohort", "Clinical incremental value and fibrinogen enhancement")
        elif group == "Reviewer 1 Major" and number == "4":
            loc += "; Additional file 2, Table S4"
        elif group == "Reviewer 2":
            loc += "; Additional file 1, Supplementary Figure S1"
        add_label(doc, "Changes in manuscript: ", f"{change_text} ({loc}).")
        coverage.append([group, number, exact_comment, "Verbatim source attachment", status, "RESPONSE_AND_LOCATION_PROVIDED"])

    doc.add_heading("Editor Comments", level=1)
    for i, text in enumerate(editor_comments, 1):
        add_comment_block("Editor", str(i), f"Editor Comment {i}", text,
                          "CLOSED_WITH_EXPLICIT_LIMITATION" if i in (1, 2, 5, 6, 10) else "CLOSED")
    add_comment_block("Editor ethics", "1", "Additional editorial ethics requirement", ethics_comment, "CLOSED")

    doc.add_heading("Reviewer 1", level=1)
    for i, text in enumerate(r1_major, 1):
        add_comment_block("Reviewer 1 Major", str(i), f"Major Comment {i}", text,
                          "CLOSED_WITH_EXPLICIT_LIMITATION" if i in (1, 4, 5, 7) else "CLOSED")
    for i, text in enumerate(r1_minor, 1):
        add_comment_block("Reviewer 1 Minor", str(i), f"Minor Comment {i}", text,
                          "CLOSED_WITH_EXPLICIT_LIMITATION" if i in (1, 2) else "CLOSED")

    doc.add_heading("Reviewer 2", level=1)
    add_comment_block("Reviewer 2", "1", "Comment 1", r2_comment, "CLOSED")
    doc.add_paragraph("Sincerely,\nThe Authors")
    doc.save(RESPONSE)

    out_csv = OUT / "05_QA/WP3_1_COMMENT_COVERAGE.csv"
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    with out_csv.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, lineterminator="\n")
        writer.writerow(["source_group", "comment_number", "verbatim_comment", "source", "coverage_status", "response_location_status"])
        writer.writerows(coverage)
    if len(coverage) != 22 or any(r[4] not in {"CLOSED", "CLOSED_WITH_EXPLICIT_LIMITATION"} for r in coverage):
        raise RuntimeError(f"Unexpected coverage result: {len(coverage)} comments")
    print(f"Created {RESPONSE} with {len(coverage)} verbatim comments")


if __name__ == "__main__":
    main()
