#!/usr/bin/env python3
"""Build the WP3.2 document-only source-lineage patch without reanalysis."""

from __future__ import annotations

import csv
import difflib
import shutil
import subprocess
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX


REPO = Path(__file__).resolve().parents[1]
WP31 = REPO / "HITS_BMC_MINOR_REVISION_WP3_1_SUBMISSION_READY"
OUT = REPO / "HITS_BMC_MINOR_REVISION_WP3_2_SUBMISSION_READY"
BASE = Path.home() / "Downloads/HITS_BMC_Manuscript_v1.2_EthicsNumber-2.docx"


def paragraph_for(doc, prefix: str):
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting {prefix!r}; found {len(matches)}")
    return matches[0]


def replace_paragraph_text(paragraph, text: str, highlighted: bool = False):
    old = paragraph.text
    first_rpr = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    for child in list(paragraph._p):
        if child.tag.endswith("}pPr"):
            continue
        paragraph._p.remove(child)
    matcher = difflib.SequenceMatcher(a=old, b=text, autojunk=False)
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            content, changed = text[j1:j2], False
        elif tag in ("insert", "replace"):
            content, changed = text[j1:j2], True
        else:
            continue
        if content:
            run = paragraph.add_run(content)
            if first_rpr is not None:
                run._r.insert(0, deepcopy(first_rpr))
            if highlighted and changed:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def set_cell_text(cell, text: str, highlighted: bool = False):
    replace_paragraph_text(cell.paragraphs[0], text, highlighted=highlighted)
    for extra in cell.paragraphs[1:]:
        replace_paragraph_text(extra, "")


def replace_embedded_figures(docx_path: Path, figure_paths: list[Path]):
    doc = Document(docx_path)
    image_paragraphs = [p for p in doc.paragraphs if p._p.xpath(".//a:blip")]
    if len(image_paragraphs) != len(figure_paths):
        raise RuntimeError(f"Expected {len(figure_paths)} embedded figures; found {len(image_paragraphs)}")
    for para, figure_path in zip(image_paragraphs, figure_paths):
        blip = para._p.xpath(".//a:blip")[0]
        rel_id = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        doc.part.rels[rel_id]._target._blob = figure_path.read_bytes()
    doc.save(docx_path)


def copy_submission_sources():
    if OUT.exists():
        raise FileExistsError(f"Refusing to overwrite existing package: {OUT}")
    for sub in ("01_MANUSCRIPT", "02_RESPONSE", "03_SUPPLEMENT", "04_FIGURES", "05_QA", "06_SUBMISSION"):
        (OUT / sub).mkdir(parents=True, exist_ok=True)
    for sub in ("source_data", "r_scripts"):
        shutil.copytree(WP31 / "04_FIGURES" / sub, OUT / "04_FIGURES" / sub)
    for path in (WP31 / "04_FIGURES").iterdir():
        if path.is_file() and not path.name.endswith(".inspect.ndjson"):
            shutil.copy2(path, OUT / "04_FIGURES" / path.name)
    shutil.copy2(WP31 / "01_MANUSCRIPT/HITS_BMC_Revised_Manuscript_CLEAN.docx",
                 OUT / "01_MANUSCRIPT/HITS_BMC_Revised_Manuscript_CLEAN.docx")
    shutil.copy2(WP31 / "03_SUPPLEMENT/Additional_file_1_Revised_Supplementary_Methods_and_Figures.docx",
                 OUT / "03_SUPPLEMENT/Additional_file_1_Revised_Supplementary_Methods_and_Figures.docx")
    shutil.copy2(WP31 / "03_SUPPLEMENT/Additional_file_2_Revised_Supplementary_Tables.xlsx",
                 OUT / "03_SUPPLEMENT/Additional_file_2_Revised_Supplementary_Tables.xlsx")
    if not BASE.exists():
        raise FileNotFoundError(f"User-confirmed original manuscript is missing: {BASE}")


def patch_manuscript(clean_path: Path):
    doc = Document(clean_path)
    patches = [
        ("Methods: This single-center retrospective study included",
         "Methods: This single-center retrospective study included 1,820 patients treated at the First Affiliated Hospital of Xinjiang Medical University from 1 January 2020 to 1 January 2026; 453 met the definite AMI definition. Core-7 combined neutrophil, lymphocyte, monocyte, platelet, mean platelet volume, red cell distribution width, and hemoglobin using elastic-net logistic regression. Performance was assessed by repeated five-fold cross-validation with preprocessing and tuning restricted to training data. PIV was the principal comparator, the Enhanced model added fibrinogen, and temporal sensitivity analysis used a six-month buffer based on the deidentified hospitalization/coronary-angiography date axis."),
        ("Results: PIV showed the highest discrimination",
         "Results: PIV showed the highest discrimination among conventional inflammatory indices (area under the receiver operating characteristic curve [AUC] 0.708, 95% confidence interval [CI] 0.678-0.735). Core-7 achieved an AUC of 0.725 (95% CI 0.697-0.753), a modest improvement over PIV (ΔAUC 0.017, 95% CI 0.003-0.031). Enhanced achieved an AUC of 0.735 (95% CI 0.710-0.763) and improved on Core-7 by 0.010 (95% CI 0.003-0.017). In the later group of the temporal sensitivity analysis, Core-7 and PIV had similar AUCs (0.701 and 0.699; ΔAUC 0.002, 95% CI -0.021 to 0.025), and Enhanced had an AUC of 0.719; this analysis was not formal temporal validation."),
        ("Conclusions: Joint modeling of routine hematologic measurements",
         "Conclusions: Joint modeling of routine hematologic measurements provided modest additional information beyond a fixed inflammatory ratio in internal evaluation. Core-7 retained moderate discrimination in the temporal sensitivity analysis, although its incremental advantage over PIV attenuated. Fibrinogen provided a small additional increment over Core-7 in internal validation. Further independent, time-anchored evaluation is warranted; these findings do not establish diagnostic or triage readiness."),
        ("This was a single-center retrospective observational study conducted",
         "This was a single-center retrospective observational study conducted in the Department of Cardiology, The First Affiliated Hospital of Xinjiang Medical University, Urumqi, Xinjiang, China, using records from Coronary Heart Disease Unit I. The study period was 1 January 2020 to 1 January 2026, and the source population comprised patients who underwent coronary angiography during this period. Because a complete eligible-patient screening log was unavailable, consecutive enrollment could not be confirmed. For the temporal sensitivity analysis, institutionally deidentified hospitalization/coronary-angiography dates were shifted on a patient-specific basis within +/-182 days. The study evaluated retrospective AMI phenotype discrimination rather than future-event prediction."),
        ("The Core-7 predictors were absolute neutrophil",
         "The Core-7 predictors were absolute neutrophil, lymphocyte, and monocyte counts, platelet count, MPV, RDW, and hemoglobin; fibrinogen was added for the Enhanced model. The hematologic measurements used for analysis were obtained in association with the target hospitalization. For patients with multiple source records, one record was selected in descending priority: non-empty discharge-diagnosis information; greater completeness of absolute neutrophil, lymphocyte, monocyte, and platelet values (0-4); presence of a parsable CBC/WBC timestamp; higher qc_nonmissing_count; and original source-row order as the final tie-break. Historical timestamp fields in the exported research dataset contained legacy registration inconsistencies and could not reliably reconstruct exact CBC sampling time relative to admission, coronary angiography, or AMI diagnosis for every patient. Accordingly, the selection rule was not a clinical-time hierarchy, and the measurements are not described as uniformly first-admission, pre-angiography, pre-diagnostic, or pre-treatment. Laboratory units were 10^9/L for leukocyte and platelet counts, fL for MPV, percent for RDW, and g/L for hemoglobin and fibrinogen. Analyzer, reagent, and platform details were not included in the deidentified research database but can be retrieved from institutional laboratory records if required during peer review. Internal CBC consistency was checked by comparing recorded absolute differential counts with values calculated from total white-cell counts and differential percentages; absolute counts were used for all derived inflammatory indices."),
        ("To comply with institutional privacy regulations, exact calendar dates",
         "The temporal sensitivity analysis was based on institutionally deidentified hospitalization/coronary-angiography dates associated with the 2020-2026 study-period cohort. These dates were shifted on a patient-specific basis within +/-182 days; a six-month buffer around the shifted-date split was excluded to reduce potential misclassification of temporal ordering. Legacy WBC_test_time fields were not used for temporal allocation. The retained analysis record identifies an earlier development group, an excluded buffer, and a later fixed-evaluation group, but does not retain a source-verified calendar cutpoint. We therefore report the groups by analysis role rather than supply the unsupported calendar cutoff from an earlier draft. The groups were development N=1,001 (AMI=196), buffer N=271 (AMI=79), and later N=548 (AMI=178). This temporal sensitivity analysis is not formal temporal validation."),
        ("Exploratory deidentified date-axis sensitivity analysis", "Temporal sensitivity analysis"),
        ("Exploratory date-axis and diagnostic sensitivity analyses", "Temporal and diagnostic sensitivity analyses"),
        ("An exploratory deidentified date-axis sensitivity analysis used the prespecified development, buffer, and later groups:",
         "The temporal sensitivity analysis used the deidentified hospitalization/coronary-angiography date axis from the 2020-2026 study period with a six-month buffer. The exact calendar cutpoint was not retained in the available analysis record, so the unsupported cutoff from an earlier draft is not reported. The groups were development N=1,001 (AMI=196), buffer N=271 (AMI=79), and later N=548 (AMI=178); the earlier group was used for development, the later group for fixed evaluation, and the buffer was excluded. In the later group, Core-7, PIV, and Enhanced AUCs were 0.701, 0.699, and 0.719, respectively. Core-7 minus PIV ΔAUC was 0.002 (95% CI -0.021 to 0.025), with an interval including zero. The incremental advantage was attenuated in this temporal sensitivity analysis, which is not formal temporal validation. The high-specificity comparison contained 138 strict-text AMI cases and 288 strict CAD/angina controls (N=426); this discharge-text rule was not independent clinical adjudication. Exploratory decision-curve results for the four clinical models are shown in Supplementary Figure S1 and do not establish clinical utility."),
        ("Figure 4. Exploratory deidentified date-axis sensitivity analysis.",
         "Figure 4. Temporal sensitivity analysis. (A) Development group, excluded buffer, and later group. (B) AUCs with 95% confidence intervals across the later-group, high-specificity discharge-text, and conservative temporal sensitivity comparisons."),
        ("For the later group in the exploratory deidentified date-axis analysis,",
         "For the later group in the temporal sensitivity analysis, Core-7 minus PIV ΔAUC was 0.002 (95% CI -0.021 to 0.025), and the interval included zero. This analysis is not formal temporal validation. Alternative date-shift estimates are presented as point estimates because confidence intervals were not available from those analyses."),
        ("The incremental advantage of Core-7 over PIV was attenuated in the exploratory deidentified date-axis sensitivity analysis.",
         "The incremental advantage of Core-7 over PIV was attenuated in the temporal sensitivity analysis. Core-7 and PIV had AUCs of approximately 0.701 and 0.699, respectively, and the paired confidence interval included zero. The temporal sensitivity analysis used deidentified hospitalization/coronary-angiography dates from the study period but is not formal temporal validation. Independent cohorts with standardized clinical time points will be needed to determine how well this performance generalizes."),
        ("Fibrinogen provided a modest additional increment over Core-7 in the primary internal analysis; an increment was also observed in the exploratory deidentified date-axis sensitivity analysis, which was not formal temporal validation.",
         "Fibrinogen provided a modest additional increment over Core-7 in the primary internal analysis; an increment was also observed in the temporal sensitivity analysis, which was not formal temporal validation. This finding is clinically plausible because fibrinogen reflects both acute-phase activation and the coagulation milieu [17]. However, fibrinogen was not available as uniformly as the CBC variables, and the Enhanced model was not consistently superior to PIV across all analyses. The inverse RDW coefficient and changes in the direction of MPV, hemoglobin, and platelet associations between crude and multivariable analyses also caution against assigning simple mechanistic interpretations to individual coefficients [13-16]. The value of the model lies in the combined hematologic pattern rather than in redefining these measurements as isolated AMI risk factors."),
        ("Several features strengthen the analysis, including predefined diagnostic mapping, preprocessing restricted to the training data during cross-validation, common cross-validated predictions for model comparisons, assessment of predictor stability, and an exploratory date-axis sensitivity analysis.",
         "Several features strengthen the analysis, including predefined diagnostic mapping, preprocessing restricted to the training data during cross-validation, common cross-validated predictions for model comparisons, assessment of predictor stability, and a temporal sensitivity analysis. The study also has important limitations. It is retrospective and single-center; AMI status was based primarily on discharge-diagnosis text rather than independent adjudication with serial biomarkers, electrocardiography, imaging, and full clinical context; and no independent external cohort was available. The source export lacked a complete eligible-patient screening log, so consecutive enrollment could not be confirmed. The extract did not contain validated index-relative prior MI or prior revascularization fields, so their prevalence and predictor roles could not be assessed. Discharge-diagnosis availability contributed to source-row ranking, and the effect of this availability-based selection cannot be quantified without the upstream extraction logic. Because admission dates were unavailable for many patients, age was derived using the available CBC/WBC timestamp as a fallback reference date and was therefore not uniformly anchored to the index angiography hospitalization. Although the analyzed hematologic measurements were associated with the target hospitalization, exact sampling times relative to admission, coronary angiography, and AMI diagnosis could not be reconstructed uniformly from the exported research dataset. Fibrinogen measurements were less complete than CBC measurements. The temporal sensitivity analysis is not formal temporal validation. These limitations define the next step: independent evaluation with a verified clinical index time, harmonized laboratory measurements, adjudicated AMI status, and prospective assessment of calibration and clinical utility [20,21,26-30]."),
        ("Multidimensional analysis of routine CBC variables provided modest additional information beyond a fixed PIV ratio in the internal evaluation. Core-7 retained moderate discrimination in the exploratory deidentified date-axis sensitivity analysis, although its incremental advantage over PIV was attenuated.",
         "Multidimensional analysis of routine CBC variables provided modest additional information beyond a fixed PIV ratio in the internal evaluation. Core-7 retained moderate discrimination in the temporal sensitivity analysis, although its incremental advantage over PIV was attenuated. Addition of fibrinogen provided a further small increment. These findings warrant independent evaluation with prospectively time-anchored laboratory measurements and standardized clinical adjudication; they do not establish diagnostic or triage readiness."),
    ]
    for prefix, text in patches:
        replace_paragraph_text(paragraph_for(doc, prefix), text)

    # Retain the original Figure 4 result values, but standardize the sensitivity label.
    label_map = {
        "Exploratory deidentified date-axis later group": "Temporal sensitivity, later group",
        "Conservative exploratory date-axis sensitivity (±12-month buffer)": "Conservative temporal sensitivity (±12-month buffer)",
    }
    table = doc.tables[2]
    for row in table.rows[1:]:
        old = row.cells[0].text
        if old in label_map:
            set_cell_text(row.cells[0], label_map[old])
    doc.save(clean_path)


def patch_supplement(path: Path):
    doc = Document(path)
    changes = [
        ("Core-7 included absolute neutrophil, lymphocyte, and monocyte counts, platelet count, MPV, RDW, and hemoglobin;",
         "Core-7 included absolute neutrophil, lymphocyte, and monocyte counts, platelet count, MPV, RDW, and hemoglobin; fibrinogen was added only in the Enhanced model. The hematologic measurements used for analysis were obtained in association with the target hospitalization. A completeness-based rule selected one record per patient by prioritizing non-empty discharge-diagnosis information, greater completeness of the four absolute counts, a parsable CBC/WBC timestamp, higher qc_nonmissing_count, and original source-row order as tie-break. Historical timestamp fields in the exported research dataset contained legacy registration inconsistencies and could not reliably reconstruct exact CBC sampling time relative to admission, coronary angiography, or AMI diagnosis for every patient; therefore, selected measurements are not described as uniformly first-admission, pre-angiography, pre-diagnostic, or pre-treatment."),
        ("An exploratory deidentified date-axis sensitivity analysis used the prespecified development, buffer, and later groups:",
         "The temporal sensitivity analysis used institutionally deidentified hospitalization/coronary-angiography dates from the 2020-2026 study period after patient-specific shifts within +/-182 days. A six-month buffer around the shifted-date split was excluded to reduce potential misclassification of temporal ordering; the anomalous legacy WBC_test_time fields were not used for temporal allocation. The retained analysis record does not contain a source-verified calendar cutpoint, so the unsupported calendar cutoff from an earlier draft is not reported. The groups were development N=1,001 (AMI=196), buffer N=271 (AMI=79), and later N=548 (AMI=178). The earlier group was used for model development and the later group for fixed evaluation; the buffer was excluded. These analyses are temporal sensitivity analyses, not formal temporal validation. A separate conservative +/-365-day sensitivity result is reported with its source confidence intervals in Table S11."),
        ("Date-axis and phenotype sensitivity", "Temporal and phenotype sensitivity"),
        ("In the later date-axis group, Core-7 AUC was 0.701, PIV 0.699, and Enhanced 0.719; Core-7 minus PIV ΔAUC was 0.002 (95% CI -0.021-0.025). The advantage attenuated and the interval crossed zero. The analysis is not formal temporal validation because available dates were not uniformly linked to the index angiography hospitalization.",
         "In the later group of the temporal sensitivity analysis, Core-7 AUC was 0.701, PIV 0.699, and Enhanced 0.719; Core-7 minus PIV ΔAUC was 0.002 (95% CI -0.021-0.025). The advantage attenuated and the interval crossed zero. The exact calendar cutpoint was not retained in the available analysis record, so the unsupported cutoff from an earlier draft is not reported. The temporal sensitivity analysis is not formal temporal validation."),
    ]
    for prefix, text in changes:
        para = next((p for p in doc.paragraphs if p.text.startswith(prefix)), None)
        if para is None:
            raise RuntimeError(f"Supplement paragraph not found: {prefix}")
        replace_paragraph_text(para, text)
    doc.save(path)


def rebuild_marked(clean_path: Path, marked_path: Path, figures: list[Path]):
    marked = Document(BASE)
    clean = Document(clean_path)
    if len(marked.paragraphs) != len(clean.paragraphs) or len(marked.tables) != len(clean.tables):
        raise RuntimeError("Original manuscript structure differs from clean manuscript")
    for source, final in zip(marked.paragraphs, clean.paragraphs):
        if source.text != final.text:
            replace_paragraph_text(source, final.text, highlighted=True)
    for source_table, final_table in zip(marked.tables, clean.tables):
        if len(source_table.rows) != len(final_table.rows):
            raise RuntimeError("Original manuscript table row count differs")
        seen = set()
        for ri, (source_row, final_row) in enumerate(zip(source_table.rows, final_table.rows)):
            if len(source_row.cells) != len(final_row.cells):
                raise RuntimeError("Original manuscript table column count differs")
            for ci, (source_cell, final_cell) in enumerate(zip(source_row.cells, final_row.cells)):
                if id(source_cell._tc) in seen:
                    continue
                seen.add(id(source_cell._tc))
                if source_cell.text != final_cell.text:
                    set_cell_text(source_cell, final_cell.text, highlighted=True)
    marked.save(marked_path)
    replace_embedded_figures(marked_path, figures)


def update_figure4_source():
    path = OUT / "04_FIGURES/source_data/Figure4_Revised_auc_source.csv"
    with path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
        fields = rows[0].keys()
    for row in rows:
        if row["analysis"] == "Exploratory date-axis sensitivity (±6-month buffer)":
            row["analysis"] = "Temporal sensitivity (±6-month buffer)"
        elif row["analysis"] == "Conservative date-axis sensitivity (±12-month buffer)":
            row["analysis"] = "Conservative temporal sensitivity (±12-month buffer)"
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main():
    copy_submission_sources()
    clean_path = OUT / "01_MANUSCRIPT/HITS_BMC_Revised_Manuscript_CLEAN.docx"
    marked_path = OUT / "01_MANUSCRIPT/HITS_BMC_Revised_Manuscript_MARKED.docx"
    supplement_path = OUT / "03_SUPPLEMENT/Additional_file_1_Revised_Supplementary_Methods_and_Figures.docx"
    patch_manuscript(clean_path)
    patch_supplement(supplement_path)
    update_figure4_source()

    rscript = shutil.which("Rscript")
    if not rscript:
        raise RuntimeError("Rscript is required to regenerate only the text-label-updated Figure 4")
    figure4_script = OUT / "04_FIGURES/r_scripts/wp3_1_plot_Figure4.R"
    subprocess.run([rscript, str(figure4_script)], check=True, cwd=OUT / "04_FIGURES")
    figures = [OUT / "04_FIGURES" / f"Figure{i}_Revised{'_OriginalStyle' if i == 1 else ''}.tiff" for i in range(1, 6)]
    replace_embedded_figures(clean_path, figures)
    rebuild_marked(clean_path, marked_path, figures)

    print(f"Created WP3.2 document inputs at {OUT}")


if __name__ == "__main__":
    main()
