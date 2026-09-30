#!/usr/bin/env python3
import argparse
import csv
import hashlib
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Inches, Pt, RGBColor
import openpyxl


ROOT = Path(__file__).resolve().parents[1]
WP1 = ROOT / "HITS_BMC_REVISION_WP1_EVIDENCE_AUDIT"
WP2 = ROOT / "HITS_BMC_MINOR_REVISION_WP2_TARGETED_ANALYSES"
OUT = ROOT / "HITS_BMC_MINOR_REVISION_WP3_FINAL"
EXPECTED_MASTER_SHA256 = "f87691f5876ece3e70b2835284bf2412ed57edaf83ee7e198df6caa89669a660"
PACKAGE_ZIP = ROOT / "HITS_BMC_MINOR_REVISION_WP3_FINAL.zip"
W = {
    "cohort": WP2 / "00_EXECUTIVE/COHORT_FLOW_FINAL.csv",
    "perf": WP2 / "03_PRIMARY_PERFORMANCE/PRIMARY_MODEL_PERFORMANCE_CANONICAL.csv",
    "deltas": WP2 / "03_PRIMARY_PERFORMANCE/PRIMARY_PAIRED_DELTAS_CANONICAL.csv",
    "clinical": WP2 / "04_CLINICAL_INCREMENTAL/CLINICAL_INCREMENTAL_PERFORMANCE_CANONICAL.csv",
    "clinical_missing": WP2 / "04_CLINICAL_INCREMENTAL/CLINICAL_COVARIATE_MISSINGNESS_FINAL.csv",
    "fbg_missing": WP2 / "05_FIBRINOGEN/FIBRINOGEN_MISSINGNESS_FINAL.csv",
    "traditional": WP2 / "06_TRADITIONAL_INDICES/TRADITIONAL_INDEX_CANONICAL_PERFORMANCE.csv",
    "spearman": WP2 / "07_COLLINEARITY/CORE7_CORRELATION_HEATMAP_SOURCE.csv",
    "stability": WP2 / "07_COLLINEARITY/VARIABLE_STABILITY_300_BOOTSTRAP.csv",
    "calibration": WP2 / "08_CALIBRATION/CALIBRATION_CANONICAL_SOURCE.csv",
    "cal_groups": WP2 / "08_CALIBRATION/CALIBRATION_GROUP_COUNTS.csv",
    "probability": WP2 / "08_CALIBRATION/PREDICTED_PROBABILITY_SUMMARY.csv",
    "dca": WP2 / "09_DCA/DCA_CANONICAL_SOURCE.csv",
    "high_flow": WP2 / "10_HIGH_SPECIFICITY/HIGH_SPECIFICITY_FLOW_FINAL.csv",
    "figure5": WP2 / "11_FIGURE5/FIGURE5_FINAL_CI_TABLE.csv",
    "date_axis": WP2 / "12_TEMPORAL_CONTEXT/DATE_AXIS_SENSITIVITY_FINAL_FACTS.csv",
    "roc": WP2 / "13_FIGURE_SOURCES/ROC_CANONICAL_SOURCE.csv",
    "table1_lock": None,
    "history": WP1 / "03_COHORT_FLOW/PRIOR_CARDIOVASCULAR_HISTORY_AVAILABILITY.csv",
    "multiplicity": WP1 / "03_COHORT_FLOW/RECORD_MULTIPLICITY.csv",
    "comments": WP1 / "00_EXECUTIVE_SUMMARY/REVIEWER_COMMENT_MASTER_MATRIX.csv",
    "response_skeleton": WP1 / "08_REVISION_PLANNING/RESPONSE_TO_REVIEWERS_SKELETON.md",
    "validation_audit": WP1 / "04_VALIDATION/VALIDATION_PIPELINE_AUDIT.md",
    "timing_audit": WP1 / "01_TIMING/CBC_INDEX_MEASUREMENT_RULE_AUDIT.md",
    "phenotype_audit": WP1 / "02_PHENOTYPE/AMI_LABEL_PROVENANCE_AUDIT.md",
    "consecutive_audit": WP1 / "03_COHORT_FLOW/CONSECUTIVE_SCREENING_AND_SELECTION_BIAS_AUDIT.md",
    "troponin_audit": WP1 / "06_COMPARATORS/TROPONIN_FEASIBILITY_AUDIT.md",
    "ethics_audit": WP1 / "07_ETHICS/ETHICS_CONSENT_HARD_GATE.md",
}


def read_csv(path):
    with Path(path).open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path, rows, fieldnames=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if fieldnames is None:
        fieldnames = list(rows[0].keys()) if rows else []
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def sha256(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def row_by(rows, **criteria):
    matches = [r for r in rows if all(str(r.get(k, "")) == str(v) for k, v in criteria.items())]
    if len(matches) != 1:
        raise ValueError(f"Expected one row for {criteria}, got {len(matches)}")
    return matches[0]


def number(row, key):
    return float(row[key])


def f3(value):
    return f"{float(value):.3f}"


def ci3(row, estimate, low, high):
    return f"{f3(row[estimate])} ({f3(row[low])}-{f3(row[high])})"


def make_dirs():
    for name in ("01_MANUSCRIPT", "02_RESPONSE", "03_SUPPLEMENT", "04_FIGURES",
                 "05_SOURCE_TRACEABILITY/source_data", "05_SOURCE_TRACEABILITY/figure_scripts",
                 "05_SOURCE_TRACEABILITY/approved_aggregate_sources",
                 "05_SOURCE_TRACEABILITY/source_inputs", "05_SOURCE_TRACEABILITY/code",
                 "06_SUBMISSION_GUIDE", "07_QA"):
        (OUT / name).mkdir(parents=True, exist_ok=True)


def build_figure_sources(age_csv):
    sd = OUT / "05_SOURCE_TRACEABILITY/source_data"
    perf = read_csv(W["perf"])
    roc = read_csv(W["roc"])
    model_map = {"PIV": "PIV", "Core-7": "Core-7", "Enhanced": "Enhanced"}
    perf_map = {model_map[r["model"]]: r for r in perf if r["analysis_id"] == "primary_core"}
    # Enhanced full-cohort performance is stored under a separate analysis id.
    perf_map["Enhanced"] = row_by(perf, analysis_id="primary_enhanced_imputed", model="Enhanced")
    roc_rows = []
    for r in roc:
        if r["model"] == "Enhanced (complete case)":
            continue
        model_name = r["model"].replace(" (full cohort)", "")
        p = perf_map[model_name]
        roc_rows.append({
            "model": model_name,
            "false_positive_rate": r["false_positive_rate"],
            "true_positive_rate": r["true_positive_rate"],
            "auc": p["AUC"],
            "auc_ci_lower": p["AUC_CI_lower"],
            "auc_ci_upper": p["AUC_CI_upper"],
            "curve_source": r.get("curve_source", ""),
        })
    write_csv(sd / "Figure2_roc_source.csv", roc_rows)

    cal = [r for r in read_csv(W["calibration"]) if
           (r["model"] == "Core-7" and r["analysis_id"] == "primary_core") or
           (r["model"] == "Enhanced (full cohort)" and r["analysis_id"] == "primary_enhanced_imputed")]
    write_csv(sd / "Figure3_calibration_source.csv", cal)

    date_rows = read_csv(W["date_axis"])
    write_csv(sd / "Figure4_date_axis_flow_source.csv", [{
        "stage": r["stage"].capitalize(), "N": r["N"], "AMI_n": r["AMI_n"],
        "Core_AUC": r["Core_AUC"], "PIV_AUC": r["PIV_AUC"],
        "Core_minus_PIV_delta_AUC": r["Core_minus_PIV_delta_AUC"],
        "delta_CI_lower": r["delta_CI_lower"], "delta_CI_upper": r["delta_CI_upper"],
        "Enhanced_AUC": r["Enhanced_AUC"], "analysis_label": r["analysis_label"],
    } for r in date_rows])
    shutil.copy2(W["figure5"], sd / "Figure5_robustness_source.csv")
    shutil.copy2(W["dca"], sd / "FigureS1_dca_source.csv")
    corr = read_csv(W["spearman"])
    write_csv(sd / "FigureS3_correlation_source.csv", [{
        "variable": r["variable_x"], "variable2": r["variable_y"], "rho": r["spearman_rho"]
    } for r in corr])
    if Path(age_csv).resolve() != (sd / "Table1_age_rebuilt_source.csv").resolve():
        shutil.copy2(age_csv, sd / "Table1_age_rebuilt_source.csv")

    flow = {r["stage"]: r for r in read_csv(W["cohort"])}
    multiplicity = {int(r["rows_per_patient"]): r for r in read_csv(W["multiplicity"])}
    nodes = [
        {"id":"source","label":"2,548 source records","x":0.8,"y":0.0,"w":1.65,"h":0.78,"group":"source"},
        {"id":"patients","label":f"2,279 unique patients\n{multiplicity[1]['patients']} with 1 row; {multiplicity[2]['patients']} with 2 rows","x":2.85,"y":0.0,"w":2.2,"h":1.02,"group":"source"},
        {"id":"dx_available","label":"1,944 with discharge diagnosis","x":4.95,"y":1.25,"w":2.15,"h":0.82,"group":"source"},
        {"id":"dx_missing","label":"335 without discharge diagnosis","x":4.95,"y":-1.25,"w":2.15,"h":0.82,"group":"exclusion"},
        {"id":"ami","label":"456 definite AMI\nbefore core-CBC filter","x":7.25,"y":2.1,"w":2.12,"h":0.9,"group":"phenotype"},
        {"id":"nonami","label":"1,372 definite non-AMI CAD\nbefore core-CBC filter","x":7.25,"y":0.68,"w":2.42,"h":0.9,"group":"phenotype"},
        {"id":"other","label":"116 other diagnoses\n5 ambiguous/review; 111 not CAD","x":7.25,"y":-0.88,"w":2.45,"h":0.92,"group":"exclusion"},
        {"id":"primary","label":"Primary cohort: N=1,820\nAMI 453; non-AMI CAD 1,367\n8 lacked required core CBC (3 + 5)","x":10.05,"y":0.78,"w":2.55,"h":1.28,"group":"primary"},
    ]
    edges = [
        {"from_id":"source","to_id":"patients"},
        {"from_id":"patients","to_id":"dx_available"},
        {"from_id":"patients","to_id":"dx_missing"},
        {"from_id":"dx_available","to_id":"ami"},
        {"from_id":"dx_available","to_id":"nonami"},
        {"from_id":"dx_available","to_id":"other"},
        {"from_id":"ami","to_id":"primary"},
        {"from_id":"nonami","to_id":"primary"},
    ]
    write_csv(sd / "Figure1_flow_nodes.csv", nodes)
    write_csv(sd / "Figure1_flow_edges.csv", edges)

    hs = {r["stage"]: r for r in read_csv(W["high_flow"])}
    hs_nodes = [
        {"id":"primary_ami","label":"Primary AMI phenotype\nN=453","x":0.95,"y":0.0,"w":1.8,"h":0.82,"group":"primary"},
        {"id":"early","label":f"Date-axis development\nN={hs['date_axis_early']['N']}","x":3.0,"y":2.05,"w":2.0,"h":0.82,"group":"source"},
        {"id":"buffer","label":f"Date-axis buffer\nN={hs['date_axis_buffer']['N']}","x":3.0,"y":0.65,"w":2.0,"h":0.82,"group":"exclusion"},
        {"id":"later","label":f"Date-axis later group\nN={hs['date_axis_later']['N']}","x":3.0,"y":-0.8,"w":2.0,"h":0.82,"group":"source"},
        {"id":"broad_controls","label":f"Later broad non-AMI controls\nN={hs['later_broad_CAD_angina_controls']['N']}","x":5.45,"y":1.5,"w":2.4,"h":0.86,"group":"source"},
        {"id":"strict_ami","label":f"Strict-text AMI\nN={hs['later_high_specificity_AMI']['N']}","x":5.45,"y":-0.25,"w":2.0,"h":0.82,"group":"phenotype"},
        {"id":"ami_not_strict","label":f"Did not meet strict text rule\nN={hs['later_AMI_not_meeting_strict_text_rule']['N']}","x":5.45,"y":-1.75,"w":2.5,"h":0.82,"group":"exclusion"},
        {"id":"strict_controls","label":f"Strict CAD/angina controls\nN={hs['later_strict_CAD_angina_controls']['N']}","x":8.05,"y":1.5,"w":2.25,"h":0.82,"group":"phenotype"},
        {"id":"controls_excluded","label":f"Excluded from strict control\nN={hs['later_broad_controls_excluded']['N']}","x":8.05,"y":-0.1,"w":2.25,"h":0.82,"group":"exclusion"},
        {"id":"comparison","label":f"High-specificity comparison\nN={hs['later_high_specificity_comparison']['N']}\nAMI 138 + controls 288","x":10.75,"y":0.15,"w":2.45,"h":1.02,"group":"primary"},
    ]
    hs_edges = [
        {"from_id":"primary_ami","to_id":"early"},
        {"from_id":"primary_ami","to_id":"buffer"},
        {"from_id":"primary_ami","to_id":"later"},
        {"from_id":"later","to_id":"strict_ami"},
        {"from_id":"later","to_id":"ami_not_strict"},
        {"from_id":"broad_controls","to_id":"strict_controls"},
        {"from_id":"broad_controls","to_id":"controls_excluded"},
        {"from_id":"strict_ami","to_id":"comparison"},
        {"from_id":"strict_controls","to_id":"comparison"},
    ]
    write_csv(sd / "FigureS2_high_specificity_flow_nodes.csv", hs_nodes)
    write_csv(sd / "FigureS2_high_specificity_flow_edges.csv", hs_edges)

    write_csv(sd / "FigureS4_domain_schema.csv", [
        {"domain":"Immune-inflammatory", "content":"Neutrophils + lymphocytes + monocytes", "group":"core", "x":1},
        {"domain":"Platelet-related", "content":"Platelet count + MPV", "group":"core", "x":1},
        {"domain":"Erythrocyte-related", "content":"RDW + hemoglobin", "group":"core", "x":1},
        {"domain":"Coagulation", "content":"Not included", "group":"core", "x":1},
        {"domain":"Immune-inflammatory", "content":"Neutrophils + lymphocytes + monocytes", "group":"enhanced", "x":2},
        {"domain":"Platelet-related", "content":"Platelet count + MPV", "group":"enhanced", "x":2},
        {"domain":"Erythrocyte-related", "content":"RDW + hemoglobin", "group":"enhanced", "x":2},
        {"domain":"Coagulation", "content":"Fibrinogen added", "group":"enhanced", "x":2},
    ])


def copy_figure_scripts():
    dest = OUT / "05_SOURCE_TRACEABILITY/figure_scripts"
    for src in sorted((ROOT / "scripts/wp3_figures").glob("*.R")):
        shutil.copy2(src, dest / src.name)
    shutil.copy2(ROOT / "scripts/wp3_rebuild_table1_age.R",
                 dest / "wp3_rebuild_table1_age.R")


def find_para(doc, prefix):
    found = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
    if len(found) != 1:
        raise ValueError(f"Expected one paragraph beginning {prefix!r}; found {len(found)}")
    return found[0]


def set_para(p, text, label=None):
    for child in list(p._p):
        if child.tag.endswith("}pPr"):
            continue
        p._p.remove(child)
    if label and text.startswith(label):
        r = p.add_run(label)
        r.bold = True
        p.add_run(text[len(label):])
    else:
        p.add_run(text)


def replace_by_prefix(doc, prefix, text, label=None):
    p = find_para(doc, prefix)
    set_para(p, text, label=label)
    return p


def set_cell(cell, text):
    p = cell.paragraphs[0]
    set_para(p, str(text))
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def set_table_rows(table, rows):
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) < len(rows):
        table.add_row()
    if table.rows[0].cells and len(table.rows[0].cells) != len(rows[0]):
        raise ValueError("Table column count changed unexpectedly")
    for row_obj, values in zip(table.rows, rows):
        if len(row_obj.cells) != len(values):
            raise ValueError(f"Table row has {len(row_obj.cells)} cells but {len(values)} values")
        for cell, value in zip(row_obj.cells, values):
            set_cell(cell, value)


def prevent_table_row_splitting(table):
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if not any(child.tag.endswith("}cantSplit") for child in tr_pr):
            from docx.oxml import OxmlElement
            tr_pr.append(OxmlElement("w:cantSplit"))


def replace_inline_figures(doc, image_paths):
    paras = [p for p in doc.paragraphs if p._p.xpath(".//a:blip")]
    shapes = list(doc.inline_shapes)
    if len(paras) != len(image_paths) or len(shapes) != len(image_paths):
        raise ValueError(f"Figure count mismatch: paragraphs={len(paras)}, shapes={len(shapes)}, replacements={len(image_paths)}")
    for p, shape, image_path in zip(paras, shapes, image_paths):
        width_inches = float(shape.width) / 914400.0
        set_para(p, "")
        p.add_run().add_picture(str(image_path), width=Inches(width_inches))


def build_main_manuscript(base_path, age_csv):
    doc = Document(base_path)
    age_rows = read_csv(age_csv)
    age = {r["group"]: r for r in age_rows}
    perf = read_csv(W["perf"])
    deltas = read_csv(W["deltas"])
    clinical = read_csv(W["clinical"])
    clinical_delta = deltas
    traditional = read_csv(W["traditional"])
    date_rows = read_csv(W["date_axis"])
    date_later = row_by(date_rows, stage="later")
    figure5 = read_csv(W["figure5"])
    fbg_missing = read_csv(W["fbg_missing"])
    miss = {r["variable"]: r for r in read_csv(W["clinical_missing"]) if r["group"] == "overall"}

    piv = row_by(perf, analysis_id="primary_core", model="PIV")
    core = row_by(perf, analysis_id="primary_core", model="Core-7")
    enhanced = row_by(perf, analysis_id="primary_enhanced_imputed", model="Enhanced")
    core_piv = row_by(deltas, comparison="Primary Core minus PIV")
    enh_core = row_by(deltas, comparison="Primary imputed Enhanced minus Core")
    fbg_cc_delta = row_by(deltas, comparison="Primary Fibrinogen complete-case Enhanced minus Core")
    clinical_cc = [r for r in clinical if r["record_type"] == "model_performance" and r["cohort_version"] == "fibrinogen_complete_case_1705"]
    clinical_rows = {
        "Clinical": row_by(clinical_cc, model="Clinical only"),
        "Clinical_Plus_PIV": row_by(clinical_cc, model="Clinical + PIV"),
        "Clinical_Plus_Core": row_by(clinical_cc, model="Clinical + Core-7"),
        "Clinical_Plus_Enhanced": row_by(clinical_cc, model="Clinical + Enhanced"),
    }

    replace_by_prefix(doc, "Methods: ", (
        "Methods: This single-center retrospective study included 1,820 patients with a discharge-diagnosis-defined "
        "AMI or non-AMI CAD phenotype. Core-7 jointly modeled seven routine CBC variables using elastic-net logistic "
        "regression. Internal performance was evaluated by stratified outer five-fold cross-validation repeated 10 "
        "times, with five-fold tuning within each training set. PIV was the prespecified primary conventional "
        "comparator; fibrinogen was assessed in an Enhanced model. All reported internal estimates used patient-level "
        "mean held-out probabilities."
    ))
    replace_by_prefix(doc, "Results: ", (
        f"Results: The primary cohort comprised 1,820 patients (453 AMI and 1,367 non-AMI CAD). PIV had an AUC of "
        f"{f3(piv['AUC'])} (95% CI {f3(piv['AUC_CI_lower'])}-{f3(piv['AUC_CI_upper'])}); Core-7 had an AUC of "
        f"{f3(core['AUC'])} ({f3(core['AUC_CI_lower'])}-{f3(core['AUC_CI_upper'])}), a modest paired increase of "
        f"{f3(core_piv['delta_auc_a_minus_b'])} ({f3(core_piv['delta_auc_ci_lower'])}-{f3(core_piv['delta_auc_ci_upper'])}). "
        f"The full-cohort Enhanced model had an AUC of {f3(enhanced['AUC'])} "
        f"({f3(enhanced['AUC_CI_lower'])}-{f3(enhanced['AUC_CI_upper'])})."
    ))
    replace_by_prefix(doc, "Conclusions: ", (
        "Conclusions: Joint modeling of routine hematologic measurements provided modest additional information for "
        "AMI phenotype discrimination beyond PIV in this single-center retrospective cohort. Fibrinogen contributed "
        "a further small increment. These findings do not establish diagnostic or triage readiness; independent "
        "cohorts with prospectively time-anchored laboratory measurements and standardized clinical adjudication are required."
    ))
    replace_by_prefix(doc, "Accordingly, we developed", (
        "Accordingly, we evaluated a multidimensional routine hematologic model for discrimination of an AMI phenotype "
        "among patients with CAD. We compared the CBC-only Core-7 model with prespecified conventional indices, "
        "particularly PIV, and assessed the additional information associated with fibrinogen. All analyses were "
        "phenotype-discrimination analyses rather than future-event prediction."
    ))
    replace_by_prefix(doc, "This was a single-center retrospective", (
        "This was a single-center retrospective observational study at the Department of Cardiology, The First "
        "Affiliated Hospital of Xinjiang Medical University, Urumqi, China. The source dataset was described as "
        "comprising patients undergoing coronary angiography from 1 January 2020 through 1 January 2026. The "
        "extraction query and complete screening ledger were not recovered, so the stated period and eligibility "
        "could not be independently machine-verified and consecutive enrollment is not claimed. The research extract "
        "contained deidentified dates; selected laboratory measurements could not be uniformly linked to the index "
        "angiography hospitalization."
    ))
    replace_by_prefix(doc, "Initially, 2,548", (
        "The source export contained 2,548 records for 2,279 unique patients. Of these, 2,010 patients had one source "
        "record and 269 had two source records. One record per patient was retained using a completeness-based "
        "selection rule rather than an earliest- or admission-first rule. Discharge-diagnosis text was available for "
        "1,944 patients; 456 met the definite AMI text phenotype, 1,372 met the definite non-AMI CAD phenotype, "
        "5 had an ambiguous MI mention, and 111 had no qualifying CAD/angina/MI phenotype. A further 335 lacked a "
        "discharge diagnosis. The required absolute neutrophil, lymphocyte, monocyte, and platelet measurements were "
        "available for 453 of 456 AMI and 1,367 of 1,372 non-AMI CAD patients, yielding the primary cohort of 1,820."
    ))
    replace_by_prefix(doc, "AMI phenotype was reconstructed", (
        "AMI phenotype was derived from discharge-diagnosis text using prespecified rules. Definite AMI required "
        "explicit current acute or subacute MI wording, including unambiguous STEMI, NSTEMI, or an acute wall-specific "
        "description. Old/prior MI, history of MI, prior PCI, CAD, stable angina, or unstable angina alone did not "
        "define AMI. MI mentions without clear acute or old/prior context and uncertainty wording were not classified "
        "as definite AMI. No independent cardiologist, ECG, imaging, or uniform troponin adjudication was available "
        "for the full cohort; the phenotype therefore represents a discharge-diagnosis-text classification, not "
        "formal adjudication under the Fourth Universal Definition [2]."
    ))
    replace_by_prefix(doc, "The cohort-entry rule required", (
        "The primary cohort required non-missing absolute neutrophil, lymphocyte, monocyte, and platelet counts. "
        "The strict/high-specificity sensitivity comparison included 138 later-group AMI records meeting the prespecified "
        "discharge-text rule and 288 strict CAD/angina controls without current MI, old/prior MI, or an ACS-only "
        "marker (N=426). The stricter AMI rule required an MI keyword plus an acute marker, excluded subacute and "
        "uncertain wording, and required a STEMI/NSTEMI/ST-segment subtype or explicit acute-MI phrase/site. These "
        "criteria were based on discharge text and use information available at or after "
        "discharge; they are not a prospective baseline rule or independent clinical adjudication."
    ))
    replace_by_prefix(doc, "The Core-7 predictors were", (
        "The Core-7 predictors were absolute neutrophil, lymphocyte, and monocyte counts, platelet count, MPV, RDW, "
        "and hemoglobin; fibrinogen was added only for the Enhanced model. Recorded units were 10^9/L for cell "
        "counts, fL for MPV, percent for RDW, and g/L for hemoglobin and fibrinogen. For patients with multiple "
        "source records, one patient-level record was selected by prioritizing a non-empty discharge diagnosis, "
        "greater completeness of absolute neutrophil, lymphocyte, monocyte, and platelet measurements, availability "
        "of a parsable WBC/CBC timestamp, higher overall non-missingness, and original source-row order as the final "
        "tie-break. The extract did not provide sufficient encounter-level identifiers or angiography/diagnosis "
        "timestamps to determine uniformly whether a selected measurement was the first admission sample or preceded "
        "angiography, diagnosis, or treatment. All 1,820 selected measurements therefore remain timing-unverified "
        "relative to the index angiography hospitalization and are interpreted as patient-level routine laboratory "
        "features associated with the discharge-diagnosis AMI phenotype."
    ))
    replace_by_prefix(doc, "Using absolute counts", (
        "Using absolute counts, we calculated NLR=Neut/Lymph, PLR=PLT/Lymph, MLR=Mono/Lymph, SII=PLT×Neut/Lymph, "
        "SIRI=Neut×Mono/Lymph, PIV=PLT×Neut×Mono/Lymph, and HRR=Hb/RDW. PIV was the prespecified primary "
        "conventional comparator; its prespecified log(1+PIV) transformation and training-fold standardization were "
        "applied in an unpenalized logistic model on the same primary cohort and outer-fold partitions as Core-7. "
        "All indices remained continuous and were not dichotomized at data-derived cutoffs [3-17]."
    ))
    replace_by_prefix(doc, "Core-7 used the seven CBC predictors", (
        "Core-7 used the seven CBC predictors, and the Enhanced model additionally included fibrinogen. Elastic-net "
        "logistic regression evaluated inverse-penalty C values of 0.10, 1.00, and 10.00 and mixing parameters "
        "of 0.25, 0.50, and 0.75 with tuning performed by five-fold "
        "cross-validation within each outer training set. Neutrophils, lymphocytes, monocytes, platelets, and "
        "fibrinogen were log(1+x) transformed and standardized; MPV, RDW, and hemoglobin were standardized on their "
        "original scale. Core-7 was complete for all 1,820 primary-cohort patients. In the full-cohort Enhanced "
        "analysis, missing fibrinogen was imputed using a median estimated within each training fold. A paired "
        "complete-case Core/Enhanced comparison was also conducted in the same 1,705 patients [18,19]."
    ))
    replace_by_prefix(doc, "Primary internal validation used", (
        "Primary internal validation used stratified outer five-fold cross-validation repeated 10 times, with "
        "five-fold tuning confined to each outer training set. Each patient received one held-out probability per "
        "repeat; the final patient-level out-of-fold probability was the arithmetic mean of the 10 held-out "
        "probabilities. Identical outer partitions were used for paired model comparisons where applicable. "
        "Imputation, transformations, centering/scaling, and hyperparameter selection were estimated in training "
        "folds only. AUC, Brier score, calibration intercept, and calibration slope were calculated from the mean "
        "out-of-fold predictions. Confidence intervals were obtained from 1,000 patient-level bootstrap resamples "
        "of the fixed mean predictions and outcomes, without refitting the full model-development pipeline; they "
        "therefore do not include all model-development uncertainty. No univariable P-value screening or stepwise "
        "selection was used [20-26]."
    ))
    replace_by_prefix(doc, "To comply with institutional privacy", (
        "An exploratory deidentified date-axis sensitivity analysis used the prespecified development, buffer, and later "
        "group allocation: development N=1,001 (AMI 196), buffer N=271 (AMI 79), and later N=548 (AMI 178). The "
        "earlier group was used for model development and the later group was evaluated as a fixed held-out "
        "date-axis group in the existing analysis. Because available dates could not be uniformly linked to the index "
        "angiography hospitalization, this analysis is not formal temporal validation and does not establish "
        "transportability."
    ))
    replace_by_prefix(doc, "High-specificity phenotype results", (
        "High-specificity phenotype and strict-control results were treated as sensitivity analyses. The high-specificity "
        "comparison comprised 138 of the 178 later-group AMI records plus 288 strict CAD/angina controls; the other "
        "40 later-group AMI records did not meet the prespecified diagnosis-text rule and should not be interpreted as "
        "adjudicated misclassifications. Fibrinogen availability and date linkage were audited; no troponin model "
        "was fitted because assay platform, ULN harmonization, serial change, and index timing were not sufficiently "
        "standardized."
    ))
    replace_by_prefix(doc, "Decision-curve analysis was deferred", (
        "An exploratory decision-curve analysis was added using the same fibrinogen-complete N=1,705 patients and "
        "mean out-of-fold predictions for Clinical, Clinical+PIV, Clinical+Core-7, and Clinical+Enhanced models, "
        "with treat-all and treat-none references over thresholds 0.05-0.50 [27]. No threshold was optimized and no "
        "clinical utility or treatment benefit is inferred. Figures were generated in R from approved aggregate "
        "analysis outputs, each with independent source data and a reproducible script [31,32]. Reporting followed "
        "STROBE, with TRIPOD+AI and PROBAST+AI used as reporting and self-audit guidance [28-30]. Detailed definitions, "
        "calibration groups, predictor stability, comparator results, and sensitivity analyses are provided in "
        "Additional files 1 and 2."
    ))
    replace_by_prefix(doc, "Of the 2,279 eligible patients", (
        "The source export contained 2,548 records and 2,279 unique patients. A total of 2,010 patients had one "
        "source record and 269 had two; one patient-level record was retained using the completeness-based "
        "rule. Among 1,944 patients with a discharge diagnosis, 456 met the definite AMI text definition, 1,372 "
        "met the definite non-AMI CAD text definition, 5 had an ambiguous MI mention, and 111 did not meet the "
        "CAD/AMI phenotype target; 335 additional patients lacked a discharge diagnosis. Eight patients with a "
        "definite AMI or non-AMI CAD phenotype lacked one or more required core CBC measurements (3 AMI and 5 "
        "non-AMI CAD), leaving 1,820 patients (453 AMI and 1,367 non-AMI CAD). The strict/high-specificity "
        "comparison contained 426 patients (138 AMI and 288 strict controls). Figure 1 summarizes record "
        "consolidation and the primary cohort."
    ))
    replace_by_prefix(doc, "Continuous variables are presented", (
        "Continuous variables are presented as median [interquartile range]. Age was available for all 1,820 "
        "patients and was derived using the primary birth date (baseline birth-date fallback), with admission date "
        "as reference when available and CBC/WBC timestamp otherwise; age was elapsed days divided by 365.2425, "
        "restricted to 18-120 years. P-values comparing definite AMI with definite non-AMI CAD were calculated "
        "using the Mann-Whitney U test for continuous variables and Pearson's chi-square test for categorical "
        "variables, with Fisher's exact test used for sparse expected cell counts. Available-case denominators for "
        "hypertension, diabetes, and fibrinogen were 1,741, 1,772, and 1,705, respectively. These P-values were "
        "descriptive and were not used for predictor selection or model development."
    ))
    replace_by_prefix(doc, "Core-7 combined immune-cell", (
        "Core-7 jointly modeled immune-inflammatory, platelet-related, and erythrocyte-related measurements; "
        "fibrinogen was added as a secondary coagulation component in the Enhanced model. These groupings organize "
        "clinical interpretation and are not treated as independent causal pathways."
    ))
    replace_by_prefix(doc, "Among the prespecified traditional", (
        f"Among the prespecified conventional indices, PIV had the highest AUC ({f3(piv['AUC'])}, 95% CI "
        f"{f3(piv['AUC_CI_lower'])}-{f3(piv['AUC_CI_upper'])}). Core-7 had an AUC of {f3(core['AUC'])} "
        f"({f3(core['AUC_CI_lower'])}-{f3(core['AUC_CI_upper'])}); its paired increase over PIV was modest "
        f"(ΔAUC {f3(core_piv['delta_auc_a_minus_b'])}, 95% CI {f3(core_piv['delta_auc_ci_lower'])}-"
        f"{f3(core_piv['delta_auc_ci_upper'])}). The full-cohort Enhanced model, with training-fold fibrinogen "
        f"imputation, had an AUC of {f3(enhanced['AUC'])} ({f3(enhanced['AUC_CI_lower'])}-"
        f"{f3(enhanced['AUC_CI_upper'])}); its paired increment over Core-7 was {f3(enh_core['delta_auc_a_minus_b'])} "
        f"(95% CI {f3(enh_core['delta_auc_ci_lower'])}-{f3(enh_core['delta_auc_ci_upper'])}). Core-7 Brier score "
        f"was {f3(core['Brier'])}, calibration intercept {f3(core['calibration_intercept'])}, and calibration "
        f"slope {f3(core['calibration_slope'])}. Figure 2 shows ROC curves and Figure 3 shows calibration from "
        f"the same mean 5-fold-by-10-repeat held-out predictions."
    ))
    replace_by_prefix(doc, "The baseline clinical model comprised", (
        f"The clinical baseline comprised age, sex, hypertension, and diabetes; smoking was unavailable. In the "
        f"same fibrinogen-complete sample (N=1,705; AMI=426), AUCs were {f3(clinical_rows['Clinical']['AUC'])} "
        f"(95% CI {f3(clinical_rows['Clinical']['AUC_CI_lower'])}-{f3(clinical_rows['Clinical']['AUC_CI_upper'])}) "
        f"for Clinical, {f3(clinical_rows['Clinical_Plus_PIV']['AUC'])} "
        f"({f3(clinical_rows['Clinical_Plus_PIV']['AUC_CI_lower'])}-{f3(clinical_rows['Clinical_Plus_PIV']['AUC_CI_upper'])}) "
        f"for Clinical+PIV, {f3(clinical_rows['Clinical_Plus_Core']['AUC'])} "
        f"({f3(clinical_rows['Clinical_Plus_Core']['AUC_CI_lower'])}-{f3(clinical_rows['Clinical_Plus_Core']['AUC_CI_upper'])}) "
        f"for Clinical+Core-7, and {f3(clinical_rows['Clinical_Plus_Enhanced']['AUC'])} "
        f"({f3(clinical_rows['Clinical_Plus_Enhanced']['AUC_CI_lower'])}-{f3(clinical_rows['Clinical_Plus_Enhanced']['AUC_CI_upper'])}) "
        f"for Clinical+Enhanced. Hypertension was missing in {miss['hypertension']['missing_n']} patients and "
        f"diabetes in {miss['diabetes']['missing_n']}; each was imputed using the training-fold median. Smoking "
        f"was not added. The full paired comparison is in Additional file 2, Table S8."
    ))
    replace_by_prefix(doc, "Because AMI prevalence varied", (
        f"The exploratory deidentified date-axis allocation included development N={date_rows[0]['N']} "
        f"(AMI={date_rows[0]['AMI_n']}), buffer N={date_rows[1]['N']} (AMI={date_rows[1]['AMI_n']}), and later "
        f"N={date_later['N']} (AMI={date_later['AMI_n']}). In the later group, Core-7 AUC was "
        f"{f3(date_later['Core_AUC'])}, PIV {f3(date_later['PIV_AUC'])}, and Enhanced "
        f"{f3(date_later['Enhanced_AUC'])}. The Core-7-minus-PIV difference was "
        f"{f3(date_later['Core_minus_PIV_delta_AUC'])} (95% CI {f3(date_later['delta_CI_lower'])}-"
        f"{f3(date_later['delta_CI_upper'])}), indicating attenuation relative to the primary internal comparison. "
        "The high-specificity comparison contained 138 AMI and 288 strict controls; the 40 later-group AMI records "
        "not meeting the stricter text criteria were not adjudicated as false cases. Because date linkage to the "
        "index angiography hospitalization was incomplete, these results are not formal temporal validation. "
        "Exploratory decision-curve analysis is shown in Supplementary Figure S1."
    ))
    replace_by_prefix(doc, "For the temporal sensitivity cohort", (
        "Sensitivity analyses are summarized in Table 3 and Figure 5. Date-axis analyses are descriptive and "
        "should not be interpreted as temporal validation. The high-specificity comparison is based on discharge "
        "diagnosis text and is not prospective adjudication."
    ))
    replace_by_prefix(doc, "In the primary stability analysis", (
        "Across 300 development bootstrap resamples, selection frequencies were 100% for neutrophils, lymphocytes, "
        "monocytes, platelets, and RDW, 99% for MPV, and 95% for hemoglobin. The maximum absolute pairwise "
        "Spearman correlation was 0.4685 between neutrophils and monocytes; no strong/extreme pairwise correlation "
        "was observed. These results are supportive stability evidence and do not justify removing variables. "
        "Additional results are provided in Additional files 1 and 2."
    ))
    replace_by_prefix(doc, "In this CAD-restricted cohort", (
        f"In this retrospective CAD cohort, Core-7 achieved an AUC of {f3(core['AUC'])} (95% CI "
        f"{f3(core['AUC_CI_lower'])}-{f3(core['AUC_CI_upper'])}), with a modest paired increment over PIV "
        f"(ΔAUC {f3(core_piv['delta_auc_a_minus_b'])}, 95% CI {f3(core_piv['delta_auc_ci_lower'])}-"
        f"{f3(core_piv['delta_auc_ci_upper'])}). This difference describes internal phenotype discrimination "
        "and does not establish diagnostic superiority."
    ))
    replace_by_prefix(doc, "The incremental advantage of Core-7", (
        f"The Core-7 advantage over PIV was attenuated in the exploratory deidentified date-axis sensitivity "
        f"analysis: AUCs were {f3(date_later['Core_AUC'])} and {f3(date_later['PIV_AUC'])}, respectively, with "
        f"a paired ΔAUC of {f3(date_later['Core_minus_PIV_delta_AUC'])} (95% CI "
        f"{f3(date_later['delta_CI_lower'])}-{f3(date_later['delta_CI_upper'])}). The interval includes zero. "
        "Available dates could not be uniformly linked to the index angiography hospitalization; the analysis is "
        "not formal temporal validation and supports no transportability claim."
    ))
    replace_by_prefix(doc, "Fibrinogen provided a reproducible", (
        f"Fibrinogen was missing for {fbg_missing[0]['missing_n']}/1,820 patients "
        f"({number(fbg_missing[0], 'missing_pct'):.2f}%), including "
        f"{fbg_missing[1]['missing_n']}/453 AMI ({number(fbg_missing[1], 'missing_pct'):.2f}%) and "
        f"{fbg_missing[2]['missing_n']}/1,367 non-AMI CAD patients "
        f"({number(fbg_missing[2], 'missing_pct'):.2f}%). In the same 1,705 fibrinogen-complete patients and "
        f"identical validation partitions, Enhanced AUC was {f3(fbg_cc_delta['auc_model_a'])} versus "
        f"{f3(fbg_cc_delta['auc_model_b'])} for Core-7 (paired ΔAUC "
        f"{f3(fbg_cc_delta['delta_auc_a_minus_b'])}, 95% CI "
        f"{f3(fbg_cc_delta['delta_auc_ci_lower'])}-{f3(fbg_cc_delta['delta_auc_ci_upper'])}), a small increment."
    ))
    replace_by_prefix(doc, "Because all seven Core-7 variables", (
        "The present study was not designed to evaluate a real-time diagnostic, triage, or treatment-decision tool. "
        "Selected CBC timing relative to the index angiography hospitalization, initial AMI diagnosis, and treatment "
        "was not uniformly established. A pooled troponin comparison was not methodologically valid because "
        "assay-platform details, ULN harmonization, serial change, and index timing were incomplete; the study "
        "therefore cannot establish incremental value beyond a standardized contemporary troponin pathway."
    ))
    replace_by_prefix(doc, "Several features strengthen", (
        "This study is retrospective and single-center. The source export did not include a complete screening "
        "register or query record, so consecutive enrollment could not be confirmed. Discharge-diagnosis availability "
        "contributed to the source-row ranking, so selection was not outcome-blind at the level of diagnosis "
        "availability; the magnitude of any resulting selection effect cannot be evaluated without the upstream "
        "extraction logic. Record consolidation used a "
        "completeness-based patient-level rule, but the retained row could not be linked uniformly to an index "
        "angiography hospitalization; selected CBC measurements therefore cannot be assumed to be admission-first, "
        "pre-diagnostic, or pre-procedural. AMI phenotype was based on discharge-diagnosis text without independent "
        "cardiologist, ECG, imaging, or standardized biomarker adjudication for all patients. AUC intervals were "
        "bootstrapped from fixed mean out-of-fold predictions and do not capture full model-development uncertainty. "
        "The exploratory deidentified date-axis analysis is not formal temporal validation, and no independent "
        "external cohort was available. Troponin assays could not be harmonized to establish incremental value "
        "beyond a standardized contemporary diagnostic pathway. These limits motivate external evaluation using "
        "prospectively time-anchored laboratories, complete source lineage, and standardized clinical adjudication."
    ))
    replace_by_prefix(doc, "Multidimensional analysis of routine CBC", (
        "In this single-center retrospective CAD cohort, joint modeling of routine hematologic measurements provided "
        "modest additional information for AMI phenotype discrimination beyond PIV, while fibrinogen contributed a "
        "further small increment. These findings are internally evaluated but should not be interpreted as evidence "
        "of diagnostic or triage readiness. Independent cohorts with prospectively time-anchored laboratory "
        "measurements and standardized clinical adjudication are required."
    ))
    replace_by_prefix(doc, "The study was approved by the Ethics Committee", (
        "The study was approved by the Ethics Committee of Xinjiang Medical University (approval no. K202602-10). "
        "Given the retrospective nature of the study, the requirement for informed consent was waived by the Ethics "
        "Committee. All methods were performed in accordance with relevant guidelines and regulations."
    ))
    replace_by_prefix(doc, "All participants provided informed consent.", "Informed consent was waived by the Ethics Committee.")
    replace_by_prefix(doc, "Consent to publish:", "Consent for publication: not applicable.")

    replace_by_prefix(doc, "Figure 1.", "Figure 1. Study population flow and completeness-based patient-level record consolidation.")
    replace_by_prefix(doc, "Figure 2.", "Figure 2. Receiver operating characteristic curves from mean 5-fold-by-10-repeat held-out predictions; AUC labels are linked to the reported primary estimates.")
    replace_by_prefix(doc, "Figure 3.", "Figure 3. Calibration by 10 quantile groups from the same mean 5-fold-by-10-repeat held-out predictions used for the reported calibration statistics.")
    replace_by_prefix(doc, "Figure 4.", "Figure 4. Exploratory deidentified date-axis sensitivity allocation; available dates were not uniformly linked to the index angiography hospitalization, so this is not formal temporal validation.")
    replace_by_prefix(doc, "Figure 5.", "Figure 5. AUC estimates with 95% confidence intervals across prespecified date-axis and high-specificity phenotype sensitivity analyses.")
    replace_by_prefix(doc, "Additional file 1. DOCX.", "Additional file 1. DOCX. Revised supplementary methods and Figures S1-S4.")
    replace_by_prefix(doc, "Additional file 2. XLSX.", "Additional file 2. XLSX. Supplementary Tables S1-S13, including updated missingness, conventional-index performance, calibration, clinical incremental comparisons, fibrinogen complete-case results, high-specificity phenotype results, and exploratory date-axis sensitivity.")

    doc.tables[0].rows[1].cells[2].text = f"{number(age['overall'], 'median'):.1f} [{number(age['overall'], 'q1'):.1f}, {number(age['overall'], 'q3'):.1f}]"
    doc.tables[0].rows[1].cells[3].text = f"{number(age['AMI'], 'median'):.1f} [{number(age['AMI'], 'q1'):.1f}, {number(age['AMI'], 'q3'):.1f}]"
    doc.tables[0].rows[1].cells[4].text = f"{number(age['non_AMI_CAD'], 'median'):.1f} [{number(age['non_AMI_CAD'], 'q1'):.1f}, {number(age['non_AMI_CAD'], 'q3'):.1f}]"
    doc.tables[0].rows[1].cells[5].text = f"{number(age['overall'], 'p_value'):.3f}"

    primary_table = [
        ["Model / comparison", "N", "AMI", "AUC (95% CI)", "ΔAUC (95% CI)", "Brier", "Calibration intercept", "Calibration slope"],
        ["PIV", piv["N"], piv["AMI_n"], ci3(piv, "AUC", "AUC_CI_lower", "AUC_CI_upper"), "—", f3(piv["Brier"]), f3(piv["calibration_intercept"]), f3(piv["calibration_slope"])],
        ["Core-7", core["N"], core["AMI_n"], ci3(core, "AUC", "AUC_CI_lower", "AUC_CI_upper"), "—", f3(core["Brier"]), f3(core["calibration_intercept"]), f3(core["calibration_slope"])],
        ["Enhanced (full cohort; training-fold fibrinogen imputation)", enhanced["N"], enhanced["AMI_n"], ci3(enhanced, "AUC", "AUC_CI_lower", "AUC_CI_upper"), "—", f3(enhanced["Brier"]), f3(enhanced["calibration_intercept"]), f3(enhanced["calibration_slope"])],
        ["Core-7 minus PIV", core_piv["n_common"], core_piv["AMI_n_common"], "—", f"{f3(core_piv['delta_auc_a_minus_b'])} ({f3(core_piv['delta_auc_ci_lower'])}-{f3(core_piv['delta_auc_ci_upper'])})", "—", "—", "—"],
        ["Enhanced minus Core-7 (full cohort)", enh_core["n_common"], enh_core["AMI_n_common"], "—", f"{f3(enh_core['delta_auc_a_minus_b'])} ({f3(enh_core['delta_auc_ci_lower'])}-{f3(enh_core['delta_auc_ci_upper'])})", "—", "—", "—"],
    ]
    set_table_rows(doc.tables[1], primary_table)
    cohort_main = row_by(figure5, analysis="Symmetric +/-182-day guard-band", model="PIV")
    date_counts = {r["stage"]: r for r in date_rows}
    # Do not include event decompositions for the conservative buffer because they are absent from the approved aggregate table.
    robust_table = [
        ["Analysis", "N", "AMI", "Controls", "PIV AUC (95% CI)", "Core-7 AUC (95% CI)", "Enhanced AUC (95% CI)"],
        ["Exploratory deidentified date-axis later group", date_later["N"], date_later["AMI_n"], str(int(date_later["N"])-int(date_later["AMI_n"])),
         "", "", ""],
        ["High-specificity phenotype comparison", "426", "138", "288", "", "", ""],
        ["Conservative +/-365-day date-axis sensitivity", "361", "Not reported", "Not reported", "", "", ""],
    ]
    for i, model in enumerate(("PIV", "Core-7", "Enhanced")):
        r = row_by(figure5, analysis="Symmetric +/-182-day guard-band", model=model)
        robust_table[1][4+i] = ci3(r, "AUC", "CI_lower", "CI_upper")
        h = row_by(figure5, analysis="High-specificity phenotype", model=model)
        robust_table[2][4+i] = ci3(h, "AUC", "CI_lower", "CI_upper")
        c = row_by(figure5, analysis="Conservative +/-365-day buffer", model=model)
        robust_table[3][4+i] = ci3(c, "AUC", "CI_lower", "CI_upper")
    set_table_rows(doc.tables[2], robust_table)
    for table in doc.tables:
        prevent_table_row_splitting(table)
    doc.save(OUT / "01_MANUSCRIPT/HITS_BMC_Revised_Manuscript_CLEAN.docx")
    replace_inline_figures(doc, [
        OUT / "04_FIGURES/Figure1_Revised.png",
        OUT / "04_FIGURES/Figure2_Revised.png",
        OUT / "04_FIGURES/Figure3_Revised.png",
        OUT / "04_FIGURES/Figure4_Revised.png",
        OUT / "04_FIGURES/Figure5_Revised.png",
    ])
    doc.save(OUT / "01_MANUSCRIPT/HITS_BMC_Revised_Manuscript_CLEAN.docx")

    marked = Document(OUT / "01_MANUSCRIPT/HITS_BMC_Revised_Manuscript_CLEAN.docx")
    modified_prefixes = (
        "Methods: ", "Results: ", "Conclusions: ", "Accordingly, we evaluated", "This was a single-center",
        "The source export contained", "AMI phenotype was derived", "The primary cohort required",
        "The Core-7 predictors were absolute", "Using absolute counts", "Core-7 jointly modeled",
        "Core-7 used the seven CBC",
        "Primary internal validation used", "An exploratory deidentified date-axis", "High-specificity phenotype",
        "An exploratory decision-curve", "The source export contained 2,548", "Continuous variables",
        "Among the prespecified conventional", "The clinical baseline comprised", "The exploratory deidentified",
        "Sensitivity analyses are summarized", "Across 300 development", "In this retrospective CAD",
        "The Core-7 advantage", "Fibrinogen was missing", "The present study was not designed",
        "This study is retrospective", "In this single-center retrospective",
        "The study was approved by the Ethics Committee",
    )
    for p in marked.paragraphs:
        if any(p.text.strip().startswith(s) for s in modified_prefixes):
            for run in p.runs:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    for table_index, table in enumerate(marked.tables):
        if table_index == 0:
            cells = [table.rows[1].cells[2], table.rows[1].cells[3],
                     table.rows[1].cells[4], table.rows[1].cells[5]]
        else:
            cells = [cell for row in table.rows[1:] for cell in row.cells]
        for cell in cells:
            if cell.text.strip():
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    marked.save(OUT / "01_MANUSCRIPT/HITS_BMC_Revised_Manuscript_MARKED.docx")


def build_supplement(supp_path):
    doc = Document(supp_path)
    perf = read_csv(W["perf"])
    deltas = read_csv(W["deltas"])
    clinical = read_csv(W["clinical"])
    missing_fbg = read_csv(W["fbg_missing"])
    date_rows = read_csv(W["date_axis"])
    later = row_by(date_rows, stage="later")
    high = read_csv(W["high_flow"])
    traditional = read_csv(W["traditional"])
    perf_cc = {r["model"]: r for r in perf if r["analysis_id"] == "primary_fbg_complete_case"}
    clinical_cc = [r for r in clinical if r["record_type"] == "model_performance" and r["cohort_version"] == "fibrinogen_complete_case_1705"]
    clin_cc = {
        "Clinical": row_by(clinical_cc, model="Clinical only"),
        "Clinical_Plus_PIV": row_by(clinical_cc, model="Clinical + PIV"),
        "Clinical_Plus_Core": row_by(clinical_cc, model="Clinical + Core-7"),
        "Clinical_Plus_Enhanced": row_by(clinical_cc, model="Clinical + Enhanced"),
    }

    replace_by_prefix(doc, "The source population comprised", (
        "The source export contained 2,548 records and 2,279 unique patients. After completeness-based "
        "patient-level consolidation, 2,010 patients had one source record and 269 had two source records. "
        "The extract did not establish that these records represented distinct admissions or identify a uniform "
        "index angiography episode. Among 1,944 patients with discharge-diagnosis text, 456 were definite AMI, "
        "1,372 were definite non-AMI CAD, 5 had an ambiguous MI mention, and 111 did not meet the CAD/AMI target; "
        "335 additional records lacked a discharge diagnosis. Required absolute neutrophil, lymphocyte, monocyte, "
        "and platelet data were available in 453 AMI and 1,367 non-AMI CAD patients."
    ))
    replace_by_prefix(doc, "AMI was defined from discharge-diagnosis", (
        "AMI was classified from discharge-diagnosis text using explicit current acute/subacute MI wording or "
        "unambiguous STEMI/NSTEMI/subtype descriptions. Old/prior MI, history of MI, prior PCI, CAD, stable angina, "
        "and unstable angina alone were not AMI-positive. Troponin and CK-MB were not used to reclassify the full "
        "cohort: assay platforms, ULNs, serial changes, and index timing were not consistently available. No "
        "independent cardiologist, ECG, imaging, or uniform biomarker adjudication was performed for all patients. "
        "The phenotype is therefore discharge-diagnosis-text based and is not formal Fourth Universal Definition "
        "adjudication. Mapping details and counts are in Additional file 2, Table S2."
    ))
    replace_by_prefix(doc, "Core-7 included absolute", (
        "Core-7 included absolute neutrophil, lymphocyte, and monocyte counts, platelet count, MPV, RDW, and "
        "hemoglobin; fibrinogen was added only in the Enhanced model. Transformations and standardization were "
        "estimated within training folds. A completeness-based rule selected one record per patient by prioritizing "
        "non-empty discharge diagnosis, completeness of the four absolute counts, a parsable WBC/CBC timestamp, "
        "overall non-missingness, and original source-row order as tie-break. Because encounter-level linkage and "
        "angiography/diagnosis timestamps were not sufficiently complete, selected CBC measurements cannot be "
        "assumed to be first, admission-first, pre-angiography, pre-diagnostic, or pre-treatment. All 1,820 "
        "measurements remain timing-unverified relative to the index angiography hospitalization."
    ))
    replace_by_prefix(doc, "Primary internal validation used", (
        "Primary internal validation used stratified outer five-fold cross-validation repeated 10 times and "
        "five-fold tuning within outer training sets. Each patient received one held-out probability per repeat; "
        "the arithmetic mean of the 10 held-out probabilities was used for patient-level evaluation. Models shared "
        "outer partitions for paired comparisons where applicable. Imputation, transformation, centering/scaling, "
        "and tuning were confined to training folds. Confidence intervals were based on 1,000 patient-level "
        "bootstrap resamples of the fixed mean out-of-fold predictions and outcomes, without refitting; intervals "
        "therefore condition on those predictions and omit full model-development uncertainty. Clinical "
        "incremental analyses used age, sex, hypertension, and diabetes; smoking was unavailable. Full definitions "
        "and results are in Additional file 2, Tables S6-S8."
    ))
    replace_by_prefix(doc, "For temporal sensitivity analysis", (
        "An exploratory deidentified date-axis sensitivity analysis used the prespecified development, buffer, and later "
        "groups: development N=1,001 (AMI=196), buffer N=271 (AMI=79), and later N=548 (AMI=178). The earlier "
        "group was used for model development and the later group was evaluated as a fixed held-out date-axis "
        "group. The selected laboratory dates could not be uniformly linked to the index angiography hospitalization; "
        "this analysis is not formal temporal validation and does not establish transportability. A separate "
        "conservative +/-365-day sensitivity result is reported with its source confidence intervals in Table S11."
    ))
    replace_by_prefix(doc, "A point-based simplification", (
        "A point-based simplification was explored during prior model development but showed inferior calibration "
        "and is not the model evaluated in the current revised results. It is retained only for transparency in "
        "Additional file 2, Table S12, and is not proposed for clinical use."
    ))
    replace_by_prefix(doc, "Figure S1.", "Figure S1. Exploratory decision-curve analysis in the fibrinogen-complete sample (N=1,705; AMI=426). Thresholds span 0.05-0.50; curves are descriptive and do not establish clinical utility.")
    replace_by_prefix(doc, "Figure S2.", "Figure S2. High-specificity phenotype comparison flow. The 138 strict-text AMI cases and 288 strict CAD/angina controls were classified from discharge-diagnosis text; this was not independent clinical adjudication.")
    replace_by_prefix(doc, "Figure S3.", "Figure S3. Pairwise Spearman correlation matrix for Core-7 predictors. Maximum absolute correlation was 0.4685 between neutrophils and monocytes.")
    replace_by_prefix(doc, "Figure S4.", "Figure S4. Analytical grouping of immune-inflammatory, platelet-related, erythrocyte-related, and coagulation measurements; groups organize interpretation and do not represent independent causal pathways.")

    ref = find_para(doc, "Supplementary figures")
    additions = [
        ("Heading 2", "Conventional indices and PIV comparison"),
        ("Normal", "All seven conventional indices were assessed using the existing repeated nested cross-validation outputs. PIV was the prespecified primary comparator; NLR, PLR, MLR, SII, SIRI, and HRR are supplementary benchmarks. Estimates and 95% confidence intervals are listed in Additional file 2, Table S4; no additional primary hypothesis tests were introduced."),
        ("Heading 2", "Clinical and fibrinogen complete-case comparisons"),
        ("Normal", f"Fibrinogen was missing in {missing_fbg[0]['missing_n']}/1,820 patients ({number(missing_fbg[0], 'missing_pct'):.2f}%), 27/453 AMI patients ({number(missing_fbg[1], 'missing_pct'):.2f}%), and 88/1,367 non-AMI CAD patients ({number(missing_fbg[2], 'missing_pct'):.2f}%). Core-7 and Enhanced complete-case models were compared in the identical 1,705 patients (426 AMI, 1,279 non-AMI CAD) and identical validation partitions; their AUCs and paired difference are reported in Tables S9-S10."),
        ("Normal", f"In this same 1,705-patient sample, AUCs were {f3(clin_cc['Clinical']['AUC'])} (95% CI {f3(clin_cc['Clinical']['AUC_CI_lower'])}-{f3(clin_cc['Clinical']['AUC_CI_upper'])}) for Clinical, {f3(clin_cc['Clinical_Plus_PIV']['AUC'])} ({f3(clin_cc['Clinical_Plus_PIV']['AUC_CI_lower'])}-{f3(clin_cc['Clinical_Plus_PIV']['AUC_CI_upper'])}) for Clinical+PIV, {f3(clin_cc['Clinical_Plus_Core']['AUC'])} ({f3(clin_cc['Clinical_Plus_Core']['AUC_CI_lower'])}-{f3(clin_cc['Clinical_Plus_Core']['AUC_CI_upper'])}) for Clinical+Core-7, and {f3(clin_cc['Clinical_Plus_Enhanced']['AUC'])} ({f3(clin_cc['Clinical_Plus_Enhanced']['AUC_CI_lower'])}-{f3(clin_cc['Clinical_Plus_Enhanced']['AUC_CI_upper'])}) for Clinical+Enhanced. Table S8 also reports Brier scores, calibration statistics, and paired ΔAUC estimates."),
        ("Heading 2", "Calibration and decision-curve analysis"),
        ("Normal", "Calibration plots use the same patient-level mean 5-fold-by-10-repeat held-out probabilities underlying the reported calibration statistics. Ten quantile groups are shown where supported by the analysis source; group-specific sample sizes and predicted-probability summaries are provided in Additional file 2, Table S6. Exploratory decision-curve analysis used the four clinical models plus treat-all and treat-none references over thresholds 0.05-0.50. It is descriptive and does not establish clinical utility, treatment benefit, or a preferred operating threshold."),
        ("Heading 2", "Date-axis and phenotype sensitivity"),
        ("Normal", f"In the later date-axis group, Core-7 AUC was {f3(later['Core_AUC'])}, PIV {f3(later['PIV_AUC'])}, and Enhanced {f3(later['Enhanced_AUC'])}; Core-7 minus PIV ΔAUC was {f3(later['Core_minus_PIV_delta_AUC'])} (95% CI {f3(later['delta_CI_lower'])}-{f3(later['delta_CI_upper'])}). The advantage attenuated and the interval crossed zero. The analysis is not formal temporal validation because available dates were not uniformly linked to the index angiography hospitalization."),
        ("Normal", f"The high-specificity comparison used {high[4]['N']} definite AMI text cases and {high[7]['N']} strict CAD/angina controls. Of 178 AMI records in the prespecified later date-axis group, 138 met the stricter text rule and 40 did not; the latter are not adjudicated misclassifications. The stricter rule required an MI keyword plus an acute marker, excluded subacute and uncertain wording, and required a STEMI/NSTEMI/ST-segment subtype or explicit acute-MI phrase/site. The criteria are based on discharge information and are not a prospective baseline rule."),
    ]
    for style, text in additions:
        ref.insert_paragraph_before(text, style=style)
    replace_inline_figures(doc, [
        OUT / "04_FIGURES/FigureS1_DCA.png",
        OUT / "04_FIGURES/FigureS2_HighSpecificityFlow.png",
        OUT / "04_FIGURES/FigureS3_Core7Correlation.png",
        OUT / "04_FIGURES/FigureS4_Domains.png",
    ])
    doc.save(OUT / "03_SUPPLEMENT/Additional_file_1_Revised_Supplementary_Methods_and_Figures.docx")


def build_response_letter(location_pages=None):
    perf = read_csv(W["perf"])
    deltas = read_csv(W["deltas"])
    clinical = read_csv(W["clinical"])
    clinical_missing = {r["variable"]: r for r in read_csv(W["clinical_missing"]) if r["group"] == "overall"}
    fbg = read_csv(W["fbg_missing"])
    traditional = read_csv(W["traditional"])
    date_axis = {r["stage"]: r for r in read_csv(W["date_axis"])}
    stability = read_csv(W["stability"])
    history = read_csv(W["history"])
    piv = row_by(perf, analysis_id="primary_core", model="PIV")
    core = row_by(perf, analysis_id="primary_core", model="Core-7")
    enhanced = row_by(perf, analysis_id="primary_enhanced_imputed", model="Enhanced")
    core_piv = row_by(deltas, comparison="Primary Core minus PIV")
    cc_fbg = row_by(deltas, comparison="Primary Fibrinogen complete-case Enhanced minus Core")
    clinical_cc = [r for r in clinical if r["record_type"] == "model_performance" and r["cohort_version"] == "fibrinogen_complete_case_1705"]
    cc_clin = {
        "Clinical": row_by(clinical_cc, model="Clinical only"),
        "Clinical_Plus_PIV": row_by(clinical_cc, model="Clinical + PIV"),
        "Clinical_Plus_Core": row_by(clinical_cc, model="Clinical + Core-7"),
        "Clinical_Plus_Enhanced": row_by(clinical_cc, model="Clinical + Enhanced"),
    }
    later = date_axis["later"]
    traditional_txt = "; ".join(
        f"{r['index']} {f3(r['AUC'])} (95% CI {f3(r['AUC_CI_lower'])}-{f3(r['AUC_CI_upper'])})"
        for r in traditional
    )
    stability_map = {r["variable"]: r for r in stability if r["bootstrap_model"] == "Continuous_HITS_Core_primary"}
    stability_txt = ", ".join(
        f"{name} {number(stability_map[name], 'selection_frequency')*100:.0f}%"
        for name in ("neut", "lymph", "mono", "plt", "mpv", "rdw", "hb")
    )
    clinical_model_order = (
        ("Clinical", "Clinical only"),
        ("Clinical+PIV", "Clinical + PIV"),
        ("Clinical+Core-7", "Clinical + Core-7"),
        ("Clinical+Enhanced", "Clinical + Enhanced"),
    )
    clinical_performance_text = "; ".join(
        f"{label}: {f3(row_by(clinical_cc, model=model)['AUC'])} "
        f"(95% CI {f3(row_by(clinical_cc, model=model)['AUC_CI_lower'])}-"
        f"{f3(row_by(clinical_cc, model=model)['AUC_CI_upper'])}), "
        f"Brier {f3(row_by(clinical_cc, model=model)['Brier'])}"
        for label, model in clinical_model_order
    )
    clinical_delta_keys = (
        "Fibrinogen-complete cohort Clinical_Plus_PIV minus Clinical",
        "Fibrinogen-complete cohort Clinical_Plus_Core minus Clinical",
        "Fibrinogen-complete cohort Clinical_Plus_Enhanced minus Clinical",
        "Fibrinogen-complete cohort Clinical_Plus_Core minus Clinical_Plus_PIV",
        "Fibrinogen-complete cohort Clinical_Plus_Enhanced minus Clinical_Plus_PIV",
        "Fibrinogen-complete cohort Clinical_Plus_Enhanced minus Clinical_Plus_Core",
    )
    clinical_deltas_text = "; ".join(
        f"{r['comparison']}: {f3(r['delta_auc_a_minus_b'])} "
        f"(95% CI {f3(r['delta_auc_ci_lower'])}-{f3(r['delta_auc_ci_upper'])})"
        for r in (row_by(deltas, comparison=key) for key in clinical_delta_keys)
    )

    comments = [
        ("E1", "CBC/fibrinogen selection and measurement timing",
         "We re-audited the source export and processing rule. A patient-level row was selected using discharge-diagnosis availability, completeness of absolute neutrophil, lymphocyte, monocyte and platelet values, a parsable CBC timestamp, overall completeness, and source-row order as the final tie-break. This was not an earliest- or admission-first rule. The extract did not permit uniform linkage to the index angiography hospitalization, admission, diagnosis, or treatment time; we therefore removed any claim of admission-first, pre-angiography, or pre-diagnostic sampling and state that all 1,820 selected measurements remain timing-unverified. Fibrinogen timing is likewise not assumed.",
         "Methods: Predictors and laboratory measurements; Results: Study population; Figure 1; Limitations."),
        ("E2", "AMI label provenance and clinical adjudication",
         "The AMI phenotype was derived from discharge-diagnosis text using explicit acute/subacute MI and subtype wording. We now state that no independent cardiologist, ECG, imaging, or uniform biomarker adjudication was performed for all patients and do not describe the phenotype as formal Fourth Universal Definition adjudication.",
         "Methods: Outcome definition; Discussion: limitations; Additional file 1; Table S2."),
        ("E3", "Reconciliation of source records and unique patients",
         "The flow now distinguishes 2,548 source records from 2,279 unique patients. Of these, 2,010 had one source record and 269 had two; we describe these as multiple source records, not proven repeat admissions. One row per patient was retained by the stated completeness-based rule.",
         "Methods: Eligibility and study population; Results: Study population; Figure 1."),
        ("E4", "Nested/repeated validation, aggregation, and bootstrap intervals",
         "We explicitly report stratified outer five-fold cross-validation repeated 10 times, five-fold inner tuning, one held-out probability per patient per repeat, arithmetic averaging of the 10 held-out probabilities, matched outer partitions for paired comparisons, and training-fold-only preprocessing and tuning. The 1,000 bootstrap intervals resample fixed mean out-of-fold predictions and outcomes without refitting; we now state that they condition on those predictions and omit full model-development uncertainty.",
         "Methods: Internal validation and uncertainty; Additional file 1; Additional file 2, Table S6."),
        ("E5", "Date-axis analysis and temporal generalizability",
         f"We renamed this analysis the exploratory deidentified date-axis sensitivity analysis. The prespecified groups were development {date_axis['development']['N']}/{date_axis['development']['AMI_n']} AMI, buffer {date_axis['buffer']['N']}/{date_axis['buffer']['AMI_n']} AMI, and later {later['N']}/{later['AMI_n']} AMI. Later-group Core-7 and PIV AUCs were {f3(later['Core_AUC'])} and {f3(later['PIV_AUC'])}; paired ΔAUC was {f3(later['Core_minus_PIV_delta_AUC'])} (95% CI {f3(later['delta_CI_lower'])}-{f3(later['delta_CI_upper'])}). The advantage attenuated. Because dates could not be uniformly linked to the index angiography hospitalization, we do not call this formal temporal validation or claim transportability.",
         "Methods: Temporal sensitivity analysis; Results; Figure 4; Discussion: limitations."),
        ("E6", "High-specificity phenotype flow",
         "The revised supplementary flow shows 453 primary AMI records, partitioned into 196 development, 79 buffer, and 178 later-group records; 138 of the 178 later AMI records met the stricter discharge-text rule, while 40 did not. This rule required an MI keyword plus an acute marker, excluded subacute and uncertain wording, and required a STEMI/NSTEMI/ST-segment subtype or explicit acute-MI phrase/site. The strict comparison used 288 CAD/angina controls, for N=426. We clarify that these were text-rule exclusions, not adjudicated misclassifications, and that discharge text is not a prospective baseline criterion.",
         "Results; Additional file 1, Figure S2; Additional file 2, Tables S10-S11."),
        ("E7", "Fibrinogen missingness and fair same-sample comparison",
         f"Fibrinogen was missing in {fbg[0]['missing_n']}/1,820 overall, {fbg[1]['missing_n']}/453 AMI, and {fbg[2]['missing_n']}/1,367 non-AMI CAD patients. Core-7 and Enhanced complete-case models used the same N=1,705 patients and the same validation partitions. AUCs were {f3(cc_fbg['auc_model_b'])} and {f3(cc_fbg['auc_model_a'])}; paired ΔAUC was {f3(cc_fbg['delta_auc_a_minus_b'])} (95% CI {f3(cc_fbg['delta_auc_ci_lower'])}-{f3(cc_fbg['delta_auc_ci_upper'])}).",
         "Results: Clinical incremental value and fibrinogen enhancement; Additional file 2, Table S9."),
        ("E8", "Clinical baseline and incremental models",
         f"The baseline uses age, sex, hypertension, and diabetes; smoking was unavailable. On the same fibrinogen-complete sample (N=1,705; AMI=426), the four-model results were {clinical_performance_text}. Paired ΔAUCs were {clinical_deltas_text}. Hypertension was missing for {clinical_missing['hypertension']['missing_n']}/1,820 ({number(clinical_missing['hypertension'], 'missing_pct'):.2f}%) and diabetes for {clinical_missing['diabetes']['missing_n']}/1,820 ({number(clinical_missing['diabetes'], 'missing_pct'):.2f}%); training-fold medians were used. Full results are in Table S8.",
         "Results: Clinical incremental value and fibrinogen enhancement; Additional file 2, Table S8."),
        ("E9", "Calibration provenance and confidence intervals",
         "Figure 3 was regenerated from the same mean 5-fold-by-10-repeat held-out predictions used for the reported calibration statistics. The supplement now gives 10 quantile-group counts and predicted-probability summaries; Figure 5 and its source table display 95% confidence intervals consistently for all nine estimates.",
         "Methods: Internal validation and uncertainty; Figure 3; Figure 5; Additional file 2, Tables S6 and S11."),
        ("E10", "Conclusion and clinical positioning",
         f"We revised the conclusion to describe only modest internal phenotype-discrimination increments. Core-7 versus PIV was ΔAUC {f3(core_piv['delta_auc_a_minus_b'])}; the later date-axis estimate was {f3(later['Core_minus_PIV_delta_AUC'])} with a confidence interval crossing zero. The fibrinogen complete-case increment was {f3(cc_fbg['delta_auc_a_minus_b'])}. We explicitly state that the findings do not establish diagnostic or triage readiness.",
         "Abstract; Discussion; Conclusions."),
        ("R1-M1", "Validation rigor and temporal generalizability",
         "We clarified the nested repeated-CV design, training-only preprocessing/tuning, arithmetic mean of 10 held-out predictions, and fixed-prediction bootstrap scope. The date-axis estimate is now labeled exploratory and not formal temporal validation; the attenuation and timing-linkage limitation are explicit.",
         "Methods: Internal validation and uncertainty; Temporal sensitivity analysis; Discussion."),
        ("R1-M2", "Core predictor collinearity and stability",
         f"The maximum absolute Core-7 Spearman correlation was 0.4685, between neutrophils and monocytes. In 300 development bootstrap resamples, selection frequencies were {stability_txt}. We describe this as no strong/extreme pairwise correlation, not absence of collinearity, and retain all prespecified variables.",
         "Results: Predictor stability; Additional file 1, Figure S3; Additional file 2, Table S7."),
        ("R1-M3", "Clinical covariate missingness and handling",
         f"Age was rebuilt using the locked birth-date/reference-date rule and was available for 1,820/1,820 patients. Hypertension was missing in {clinical_missing['hypertension']['missing_n']}/1,820 ({number(clinical_missing['hypertension'], 'missing_pct'):.2f}%) and diabetes in {clinical_missing['diabetes']['missing_n']}/1,820 ({number(clinical_missing['diabetes'], 'missing_pct'):.2f}%); their medians were estimated within training folds. Smoking was unavailable and was not added.",
         "Methods: Age derivation and internal validation; Table 1; Additional file 2, Table S3."),
        ("R1-M4", "Traditional-index benchmarks",
         f"The supplement reports the seven existing repeated nested cross-validation estimates with 95% CIs: {traditional_txt}. PIV remains the prespecified primary comparator; these do not represent seven new primary tests.",
         "Methods: Conventional indices; Additional file 2, Table S4."),
        ("R1-M5", "Multiple source records and prior CAD/MI/revascularization history",
         "We distinguish 2,548 source records, 2,279 unique patients, 2,010 patients with one row, and 269 with two source rows. The available history fields did not have validated index-relative semantics, and no validated dedicated prior CAD, prior MI, prior PCI, or CABG field was recovered. We therefore do not invent or report these as verified prior-history covariates.",
         "Methods: Eligibility and study population; Results: Study population; Additional file 1."),
        ("R1-M6", "Fairness of the PIV comparison",
         "PIV was prespecified as the principal conventional comparator and used the log(1+PIV) transformation, training-fold standardization, the same primary cohort, and the same outer partitions as Core-7. Paired comparisons use the same patient-level mean held-out predictions; the manuscript describes phenotype discrimination rather than clinical superiority.",
         "Methods: Conventional inflammatory indices; Internal validation and uncertainty; Results."),
        ("R1-M7", "Troponin comparison",
         "We did not fit a troponin model or pool assay results. The recovered data lacked sufficiently complete assay-platform/manufacturer details, ULN harmonization, validated serial change, and index timing; CK-MB was present only sparsely. We acknowledge that incremental value beyond a standardized contemporary troponin pathway cannot be established.",
         "Methods: Outcome definition; Discussion: limitations."),
        ("R1-m1", "Consecutive screening and study period",
         "The source dataset was described as covering patients undergoing coronary angiography from 1 January 2020 through 1 January 2026, but the extraction query and complete eligible-patient screening ledger were not recovered. We therefore do not claim consecutive enrollment and state that the date range/eligibility could not be independently machine-verified.",
         "Methods: Study design and setting; Discussion: limitations."),
        ("R1-m2", "CBC timing clarification",
         "We report the exact completeness-based source-row selection hierarchy and explicitly state that CBC-to-index-hospitalization timing could not be reconstructed uniformly. We make no claim of first admission, pre-angiography, pre-diagnostic, or pre-treatment sampling.",
         "Methods: Predictors and laboratory measurements; Figure 1; Discussion."),
        ("R1-m3", "Informed consent for retrospective research",
         "The inconsistent statement that all participants provided informed consent was removed. The Declarations now state ethics approval K202602-10 and that the Ethics Committee waived informed consent for this retrospective study, with compliance with relevant guidelines and regulations.",
         "Declarations: Ethics approval and Consent to participate."),
        ("R2-1", "Decision-curve analysis for clinical incremental models",
         "Exploratory DCA was added in Supplementary Figure S1 using the same N=1,705 patients (AMI=426), four clinical models, treat-all/treat-none, and thresholds 0.05-0.50. No threshold was optimized. We describe net-benefit differences as exploratory and do not claim clinical utility or treatment benefit.",
         "Results: exploratory sensitivity analyses; Additional file 1, Figure S1; Additional file 2, Table S8."),
    ]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ("Heading 1", "Heading 2"):
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(12)
    title_run = title.add_run("Response to the Editor and Reviewers")
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    p = doc.add_paragraph("Submission: HITS / BMC Cardiovascular Disorders")
    p.runs[0].bold = True
    doc.add_paragraph("Submission ID: fa1e4873-011b-4f15-8b0d-29f6df21fabe")
    note = doc.add_paragraph()
    r = note.add_run("Draft limitation: ")
    r.bold = True
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    note.add_run("The exact editor letter, Reviewer 1 and Reviewer 2 reports, and portal-integrated manuscript snapshot were not recovered. The items below are explicitly non-verbatim summaries from the adopted audit instructions and are not quotations. This draft requires comparison with the original reports before submission.")
    doc.add_paragraph("We thank the Editor and reviewers for their careful assessment. We revised the manuscript within the scope of the approved analyses and have not added new model searches, biomarkers, phenotypes, datasets, or algorithms.")
    location_pages = location_pages or {}
    for cid, summary, response, location in comments:
        doc.add_heading(cid, level=2)
        p = doc.add_paragraph()
        r = p.add_run("Comment summary (non-verbatim): ")
        r.bold = True
        p.add_run(summary)
        p = doc.add_paragraph()
        r = p.add_run("Response: ")
        r.bold = True
        p.add_run(response)
        p = doc.add_paragraph()
        r = p.add_run("Changes in the revised files: ")
        r.bold = True
        page_note = location_pages.get(cid, "")
        p.add_run((page_note + "; " if page_note else "") + location)
    doc.add_heading("Ethics clarification (contextual audit item; not counted among Editor comments 1-10)", level=2)
    doc.add_paragraph("The Ethics approval and consent declarations now give approval number K202602-10, the retrospective informed-consent waiver, and compliance with relevant guidelines and regulations. Approval documents were not included in this public-safe package.")
    doc.save(OUT / "02_RESPONSE/HITS_BMC_Point_by_Point_Response.docx")

    matrix = []
    for cid, summary, response, location in comments:
        reviewer = "Editor" if cid.startswith("E") else ("Reviewer 1" if cid.startswith("R1") else "Reviewer 2")
        page_note = location_pages.get(cid, "")
        full_location = (page_note + "; " if page_note else "") + location
        matrix.append({
            "comment_id": cid,
            "comment_summary": summary,
            "response_complete": "YES_FOR_RECOVERED_SUMMARY_ONLY",
            "manuscript_change": full_location,
            "supplement_change": "Updated where relevant; see response and supplementary cross-references",
            "figure_change": "Updated where relevant using R and frozen aggregate sources",
            "new_analysis_used": "NO; frozen approved WP2 aggregate outputs only",
            "location": full_location,
            "unresolved_limitation": "Exact verbatim editor/reviewer text not recovered; response-to-source alignment cannot be certified",
            "final_status": "CLOSED_WITH_EXPLICIT_LIMITATION",
            "reviewer_group": reviewer,
        })
    matrix.append({
        "comment_id": "E11_CONTEXTUAL_NOT_COUNTED",
        "comment_summary": "Human research guidelines statement, a contextual WP1 audit item rather than one of Editor comments 1-10",
        "response_complete": "YES_FOR_RECOVERED_SUMMARY_ONLY",
        "manuscript_change": "Declarations: Ethics approval and Consent to participate",
        "supplement_change": "Not applicable",
        "figure_change": "Not applicable",
        "new_analysis_used": "NO",
        "location": "Ethics clarification",
        "unresolved_limitation": "Exact editor-report wording not recovered; approval documents were not independently inspected",
        "final_status": "CLOSED_WITH_EXPLICIT_LIMITATION",
        "reviewer_group": "Contextual",
    })
    write_csv(OUT / "05_SOURCE_TRACEABILITY/WP3_COMMENT_COVERAGE_MATRIX.csv", matrix)


def append_run_log(message):
    path = OUT / "05_SOURCE_TRACEABILITY/run.log"
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as f:
        f.write(f"{datetime.now(timezone.utc).isoformat()}Z\t{message}\n")


def sanitize_public_paths():
    text_suffixes = {".csv", ".json", ".log", ".md", ".mjs", ".py", ".r", ".txt"}
    replacements = ((str(ROOT), "$WP3_ROOT"), (str(Path.home()), "$HOME"))
    for path in OUT.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in text_suffixes:
            continue
        try:
            original = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            continue
        sanitized = original
        for source, replacement in replacements:
            sanitized = sanitized.replace(source, replacement)
        if sanitized != original:
            path.write_text(sanitized, encoding="utf-8")


def copy_source(src, dest):
    src, dest = Path(src), Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dest)
    return dest


def copy_approved_aggregate_sources():
    keys = (
        "cohort", "perf", "deltas", "clinical", "clinical_missing", "fbg_missing",
        "traditional", "spearman", "stability", "calibration", "cal_groups",
        "probability", "dca", "high_flow", "figure5", "date_axis", "roc",
        "multiplicity", "comments", "validation_audit", "timing_audit",
        "phenotype_audit", "consecutive_audit", "troponin_audit", "ethics_audit",
    )
    copied = []
    for key in keys:
        src = Path(W[key])
        if not src.is_file():
            raise FileNotFoundError(f"Approved source missing: {src.name}")
        dest = OUT / "05_SOURCE_TRACEABILITY/approved_aggregate_sources" / src.parent.name / src.name
        copy_source(src, dest)
        copied.append((key, src, dest))
    return copied


def write_source_manifest(args, copied_sources, age_csv):
    rows = []

    def add(name, digest, role, version, authoritative, status="RECOVERED"):
        rows.append({"filename": name, "SHA256": digest, "role": role, "version": version,
                     "authoritative": "YES" if authoritative else "NO", "status": status})

    for key, src, dest in copied_sources:
        add(dest.relative_to(OUT).as_posix(), sha256(dest), f"Approved WP1/WP2 source: {key}",
            "review/audit output", key != "comments")
    inputs = (
        (args.manuscript, "working manuscript input; portal snapshot not verified", "v1.2 copy", False),
        (args.supplement, "recovered supplementary-methods input; portal snapshot not verified", "recovered package", False),
        (args.workbook, "recovered supplementary-workbook input", "recovered package", False),
        (args.table1_source, "Table 1 aggregate source lock", "scientific freeze", True),
    )
    for path, role, version, authoritative in inputs:
        src = Path(path)
        dest = copy_source(src, OUT / "05_SOURCE_TRACEABILITY/source_inputs" / src.name)
        add(dest.relative_to(OUT).as_posix(), sha256(dest), role, version, authoritative)
    master_hash = sha256(args.restricted_master)
    add("RESTRICTED_MASTER_NOT_INCLUDED.csv", master_hash,
        "Restricted source used only for aggregate age reconstruction; raw file excluded",
        "audited master", True, "HASH_RECORDED_RAW_FILE_EXCLUDED")
    for name, role in (
        ("Exact_BMC_editor_decision_letter", "Exact editorial decision source"),
        ("Exact_BMC_Reviewer_1_report", "Exact reviewer report"),
        ("Exact_BMC_Reviewer_2_report", "Exact reviewer report"),
        ("Portal_integrated_submitted_manuscript", "Exact manuscript reviewed by the journal"),
    ):
        add(name, "", role, "not recovered", False, "NOT_RECOVERED")
    add("05_SOURCE_TRACEABILITY/source_data/Table1_age_rebuilt_source.csv", sha256(age_csv),
        "Aggregate age summary rebuilt from restricted source using locked definition", "WP2 age lock", True)
    write_csv(OUT / "05_SOURCE_TRACEABILITY/WP3_SOURCE_MANIFEST.csv", rows)
    return rows


def write_figure_source_map():
    figures = [
        ("Figure1_Revised", "Figure1_flow_nodes.csv; Figure1_flow_edges.csv", "plot_Figure1_Flow.R", "cohort and multiplicity aggregates"),
        ("Figure2_Revised", "Figure2_roc_source.csv", "plot_Figure2_ROC.R", "canonical OOF ROC and AUC"),
        ("Figure3_Revised", "Figure3_calibration_source.csv", "plot_Figure3_Calibration.R", "canonical mean 5x10 OOF calibration"),
        ("Figure4_Revised", "Figure4_date_axis_flow_source.csv", "plot_Figure4_DateAxis.R", "approved date-axis group counts"),
        ("Figure5_Revised", "Figure5_robustness_source.csv", "plot_Figure5_Robustness.R", "approved AUC/CI rows"),
        ("FigureS1_DCA", "FigureS1_dca_source.csv", "plot_FigureS1_DCA.R", "exploratory DCA aggregate"),
        ("FigureS2_HighSpecificityFlow", "FigureS2_high_specificity_flow_nodes.csv; FigureS2_high_specificity_flow_edges.csv", "plot_FigureS2_HighSpecificityFlow.R", "approved diagnosis-text flow counts"),
        ("FigureS3_Core7Correlation", "FigureS3_correlation_source.csv", "plot_FigureS3_Correlation.R", "approved Spearman matrix"),
        ("FigureS4_Domains", "FigureS4_domain_schema.csv", "plot_FigureS4_Domains.R", "analytic domain schema; no outcome result"),
    ]
    rows = [{"figure_stem": stem, "source_data_csv": source, "R_script": script,
             "result_source": note, "exports": "PDF; TIFF 600 dpi; PNG 600 dpi"}
            for stem, source, script, note in figures]
    write_csv(OUT / "05_SOURCE_TRACEABILITY/FIGURE_SOURCE_MAP.csv", rows)


def capture_runtime_info():
    py = subprocess.run([sys.executable, "--version"], capture_output=True, text=True)
    (OUT / "05_SOURCE_TRACEABILITY/python_runtime.txt").write_text(
        f"Python executable: {sys.executable}\n{py.stdout}{py.stderr}\nPlatform: {platform.platform()}\n",
        encoding="utf-8")
    rscript = shutil.which("Rscript")
    if not rscript:
        raise RuntimeError("Rscript is required.")
    info = subprocess.run([rscript, "-e", "cat(capture.output(sessionInfo()), sep='\\n')"],
                          capture_output=True, text=True, check=True)
    (OUT / "05_SOURCE_TRACEABILITY/R_sessionInfo.txt").write_text(info.stdout, encoding="utf-8")
    node = shutil.which("node")
    version = subprocess.run([node, "--version"], capture_output=True, text=True, check=True).stdout if node else "NOT_AVAILABLE\n"
    (OUT / "05_SOURCE_TRACEABILITY/Node_runtime.txt").write_text(version, encoding="utf-8")


def prepare_package(args):
    if OUT.exists() and any(OUT.iterdir()):
        log = OUT / "05_SOURCE_TRACEABILITY/run.log"
        if not log.is_file() or f"restricted_source_sha256={EXPECTED_MASTER_SHA256}" not in log.read_text(encoding="utf-8"):
            raise FileExistsError(f"Refusing to overwrite unrelated non-empty output: {OUT}")
    for key in ("manuscript", "supplement", "workbook", "table1_source", "restricted_master"):
        if not Path(getattr(args, key)).is_file():
            raise FileNotFoundError(f"Input missing: {Path(getattr(args, key)).name}")
    master_hash = sha256(args.restricted_master)
    if master_hash != args.restricted_master_sha256 or master_hash != EXPECTED_MASTER_SHA256:
        raise ValueError("Restricted master SHA-256 differs from the approved source lock.")
    W["table1_lock"] = Path(args.table1_source)
    make_dirs()
    copy_source(args.workbook, OUT / "03_SUPPLEMENT/Additional_file_2_Revised_Supplementary_Tables.xlsx")
    age_csv = OUT / "05_SOURCE_TRACEABILITY/source_data/Table1_age_rebuilt_source.csv"
    rscript = shutil.which("Rscript")
    if not rscript:
        raise RuntimeError("Rscript is required for locked age reconstruction.")
    cmd = [rscript, str(ROOT / "scripts/wp3_rebuild_table1_age.R"), str(args.restricted_master), str(age_csv)]
    result = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True)
    append_run_log(f"AGE_BUILD restricted_source_sha256={master_hash}; output={age_csv.relative_to(OUT).as_posix()}")
    append_run_log((result.stdout + result.stderr).strip())
    if result.returncode:
        raise RuntimeError("Age rebuild failed; see run.log.")
    copied = copy_approved_aggregate_sources()
    copy_figure_scripts()
    copy_source(ROOT / "scripts/build_wp3_package.py", OUT / "05_SOURCE_TRACEABILITY/code/build_wp3_package.py")
    copy_source(ROOT / "scripts/update_wp3_supplement_workbook.mjs", OUT / "05_SOURCE_TRACEABILITY/code/update_wp3_supplement_workbook.mjs")
    build_figure_sources(age_csv)
    write_figure_source_map()
    manifest = write_source_manifest(args, copied, age_csv)
    capture_runtime_info()
    metadata = {
        "package": "HITS_BMC_MINOR_REVISION_WP3_FINAL",
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "restricted_master_sha256": master_hash,
        "manifest_rows": len(manifest),
        "exact_review_reports_recovered": False,
        "portal_integrated_manuscript_recovered": False,
        "raw_patient_level_data_included": False,
    }
    (OUT / "05_SOURCE_TRACEABILITY/build_metadata.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    append_run_log("PREPARE_PASS; restricted row-level data excluded")
    print(json.dumps({"phase": "prepare", "out": str(OUT), "manifest_rows": len(manifest), "age_csv": str(age_csv)}))


def assemble_package(args):
    if not (OUT / "05_SOURCE_TRACEABILITY/source_data/Table1_age_rebuilt_source.csv").is_file():
        raise FileNotFoundError("Run phase=prepare first.")
    rscript = shutil.which("Rscript")
    if not rscript:
        raise RuntimeError("Rscript is required for figure rendering.")
    script_dir = OUT / "05_SOURCE_TRACEABILITY/figure_scripts"
    for script in sorted(script_dir.glob("plot_*.R")):
        cmd = [rscript, str(script), str(OUT)]
        result = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True)
        append_run_log("R_FIGURE " + " ".join(cmd))
        append_run_log((result.stdout + result.stderr).strip())
        if result.returncode:
            raise RuntimeError(f"Figure script failed: {script.name}; see run.log.")
    age_csv = OUT / "05_SOURCE_TRACEABILITY/source_data/Table1_age_rebuilt_source.csv"
    build_main_manuscript(args.manuscript, age_csv)
    build_supplement(args.supplement)
    append_run_log("ASSEMBLE_PASS; main and supplementary DOCX built from approved aggregates")
    print(json.dumps({"phase": "assemble", "manuscript": str(OUT / "01_MANUSCRIPT/HITS_BMC_Revised_Manuscript_CLEAN.docx"),
                      "supplement": str(OUT / "03_SUPPLEMENT/Additional_file_1_Revised_Supplementary_Methods_and_Figures.docx")}))


def read_pdf_pages(path):
    exe = shutil.which("pdftotext")
    if exe:
        result = subprocess.run([exe, "-layout", str(path), "-"], capture_output=True, text=True, check=True)
        return result.stdout.split("\f")
    try:
        from pypdf import PdfReader
        return [page.extract_text() or "" for page in PdfReader(str(path)).pages]
    except Exception as e:
        raise RuntimeError("pdftotext or pypdf is needed for page mapping.") from e


def find_pages(pages, terms):
    result = []
    for i, page in enumerate(pages, start=1):
        if any(term.casefold() in page.casefold() for term in terms):
            result.append(i)
    return result


def build_comment_page_map(main_pdf, supplement_pdf):
    main_pages = read_pdf_pages(main_pdf)
    supp_pages = read_pdf_pages(supplement_pdf)
    anchors = {
        "E1": (["Predictors and laboratory measurements", "Study population", "This study is retrospective"], ["CBC timing", "Figure S2."]),
        "E2": (["Outcome definition", "AMI phenotype was based on discharge-diagnosis text"], ["AMI was classified", "Figure S2."]),
        "E3": (["Eligibility and study population", "Figure 1."], ["The source population comprised"]),
        "E4": (["Internal validation and uncertainty"], ["Primary internal validation"]),
        "E5": (["Temporal sensitivity analysis", "Figure 4."], ["Date-axis and phenotype sensitivity"]),
        "E6": (["Primary and strict-control cohorts"], ["Figure S2.", "High-specificity comparison flow"]),
        "E7": (["Fibrinogen enhancement"], ["Fibrinogen was missing"]),
        "E8": (["Clinical incremental value and fibrinogen enhancement"], ["Same-sample clinical incremental comparison"]),
        "E9": (["Figure 3.", "Figure 5."], ["Calibration and decision-curve analysis"]),
        "E10": (["Abstract", "Conclusions"], ["Date-axis and phenotype sensitivity"]),
        "R1-M1": (["Internal validation and uncertainty", "Temporal sensitivity analysis"], ["Primary internal validation"]),
        "R1-M2": (["Predictor stability", "Figure S3."], ["Core-7 pairwise Spearman correlations"]),
        "R1-M3": (["Continuous variables are presented", "Hypertension", "Diabetes"], ["Missingness in the primary cohort"]),
        "R1-M4": (["Conventional inflammatory indices"], ["Conventional index performance"]),
        "R1-M5": (["Eligibility and study population"], ["Source records and prior cardiovascular history"]),
        "R1-M6": (["Conventional inflammatory indices", "Internal validation and uncertainty"], ["PIV comparison"]),
        "R1-M7": (["The present study was not designed", "Discussion"], ["AMI was classified"]),
        "R1-m1": (["Study design and setting", "This study is retrospective"], []),
        "R1-m2": (["Predictors and laboratory measurements", "Figure 1."], ["CBC timing"]),
        "R1-m3": (["Ethics approval", "Consent to participate"], []),
        "R2-1": (["Decision-curve analysis", "Figure S1."], ["Exploratory decision-curve analysis"]),
    }
    result = {}
    report = []
    for cid, (main_terms, supp_terms) in anchors.items():
        mp = find_pages(main_pages, main_terms)
        sp = find_pages(supp_pages, supp_terms) if supp_terms else []
        pieces = []
        if mp:
            pieces.append("Revised manuscript pp. " + ", ".join(map(str, mp)))
        if sp:
            pieces.append("Additional file 1 pp. " + ", ".join(map(str, sp)))
        result[cid] = "; ".join(pieces)
        report.append({"comment_id": cid, "manuscript_pages": ";".join(map(str, mp)),
                       "supplement_pages": ";".join(map(str, sp)),
                       "locator_status": "MAPPED" if (mp or sp) else "SECTION_ONLY_REVIEW_REQUIRED"})
    write_csv(OUT / "05_SOURCE_TRACEABILITY/WP3_MANUSCRIPT_PAGE_MAP.csv", report)
    return result


def response_package(args):
    for path in (args.manuscript_pdf, args.supplement_pdf):
        if not Path(path).is_file():
            raise FileNotFoundError(f"Rendered PDF missing: {Path(path).name}")
    locations = build_comment_page_map(args.manuscript_pdf, args.supplement_pdf)
    build_response_letter(locations)
    append_run_log("RESPONSE_PASS; non-verbatim comment summaries include rendered page locators")
    print(json.dumps({"phase": "response", "path": str(OUT / "02_RESPONSE/HITS_BMC_Point_by_Point_Response.docx"),
                      "mapped_comments": len(locations)}))


def extract_docx_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def run_claim_reference_audit(manuscript):
    doc = Document(manuscript)
    text = extract_docx_text(manuscript)
    terms = ("predict", "prediction", "predictive", "diagnostic", "diagnosis", "triage",
             "clinical utility", "clinical application", "validation", "temporal validation",
             "external validation", "pre-angiography", "pre-diagnostic", "admission CBC",
             "consecutive", "all participants provided informed consent")
    claim_rows = []
    for term in terms:
        matches = list(re.finditer(re.escape(term), text, flags=re.IGNORECASE))
        for match in matches:
            excerpt = re.sub(r"\s+", " ", text[max(0, match.start()-100):min(len(text), match.end()+140)])
            claim_rows.append({"term": term, "occurrence_n": len(matches), "excerpt": excerpt,
                               "review_status": "CONTEXT_REVIEW_REQUIRED"})
    write_csv(OUT / "07_QA/WP3_CLAIM_AUDIT.csv", claim_rows)

    ref_start = False
    references = {}
    for p in doc.paragraphs:
        t = p.text.strip()
        if t == "References":
            ref_start = True
            continue
        m = re.match(r"^(\d+)\.\s+", t)
        if ref_start and m:
            references[int(m.group(1))] = t
    cited = set()
    paragraph_text = "\n".join(p.text for p in doc.paragraphs)
    for match in re.finditer(r"\[([0-9,\-– ]+)\]", paragraph_text):
        for token in re.split(r",\s*", match.group(1)):
            rng = re.fullmatch(r"(\d+)\s*[-–]\s*(\d+)", token.strip())
            if rng:
                cited.update(range(int(rng.group(1)), int(rng.group(2))+1))
            elif token.strip().isdigit():
                cited.add(int(token.strip()))
    write_csv(OUT / "07_QA/WP3_CITATION_REFERENCE_AUDIT.csv", [
        {"reference_number": n, "present_in_reference_list": "YES" if n in references else "NO",
         "cited_in_text": "YES" if n in cited else "NO", "reference_entry": references.get(n, "")}
        for n in sorted(set(references) | cited)
    ])
    return text, references, cited


def build_numerical_crosscheck(manuscript, supplement):
    text = extract_docx_text(manuscript) + "\n" + extract_docx_text(supplement)
    workbook_path = OUT / "03_SUPPLEMENT/Additional_file_2_Revised_Supplementary_Tables.xlsx"
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    workbook_text = "\n".join(
        str(value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows(values_only=True)
        for value in row
        if value is not None
    )
    workbook.close()
    perf = read_csv(W["perf"])
    deltas = read_csv(W["deltas"])
    figure5 = read_csv(W["figure5"])
    age_rows = {r["group"]: r for r in read_csv(OUT / "05_SOURCE_TRACEABILITY/source_data/Table1_age_rebuilt_source.csv")}
    rows = []

    def add(location, item, source, source_value, display, evidence_text=None):
        evidence = text if evidence_text is None else evidence_text
        rows.append({"manuscript_location": location, "variable_result": item, "source_file": source,
                     "source_value": source_value, "manuscript_value": display,
                     "match": "YES" if display.casefold() in evidence.casefold() else "NO"})

    for model in ("PIV", "Core-7"):
        r = row_by(perf, analysis_id="primary_core", model=model)
        add("Table 2 / Results", f"{model} AUC (95% CI)", W["perf"].relative_to(ROOT).as_posix(),
            f"{r['AUC']} [{r['AUC_CI_lower']},{r['AUC_CI_upper']}]", ci3(r, "AUC", "AUC_CI_lower", "AUC_CI_upper"))
    enh = row_by(perf, analysis_id="primary_enhanced_imputed", model="Enhanced")
    add("Table 2 / Results", "Enhanced full cohort AUC (95% CI)", W["perf"].relative_to(ROOT).as_posix(),
        f"{enh['AUC']} [{enh['AUC_CI_lower']},{enh['AUC_CI_upper']}]", ci3(enh, "AUC", "AUC_CI_lower", "AUC_CI_upper"))
    d = row_by(deltas, comparison="Primary Core minus PIV")
    add("Table 2 / Results", "Core-7 minus PIV paired delta AUC", W["deltas"].relative_to(ROOT).as_posix(),
        f"{d['delta_auc_a_minus_b']} [{d['delta_auc_ci_lower']},{d['delta_auc_ci_upper']}]",
        f"{f3(d['delta_auc_a_minus_b'])} ({f3(d['delta_auc_ci_lower'])}-{f3(d['delta_auc_ci_upper'])})")
    core = row_by(perf, analysis_id="primary_core", model="Core-7")
    for field in ("Brier", "calibration_intercept", "calibration_slope"):
        add("Table 2", field, W["perf"].relative_to(ROOT).as_posix(), core[field], f3(core[field]))
    d = row_by(deltas, comparison="Primary Fibrinogen complete-case Enhanced minus Core")
    add("Results / Table S9", "Enhanced minus Core-7 complete-case delta AUC", W["deltas"].relative_to(ROOT).as_posix(),
        f"{d['delta_auc_a_minus_b']} [{d['delta_auc_ci_lower']},{d['delta_auc_ci_upper']}]",
        f"{f3(d['delta_auc_a_minus_b'])} ({f3(d['delta_auc_ci_lower'])}-{f3(d['delta_auc_ci_upper'])})",
        evidence_text=workbook_text)
    for group in ("overall", "AMI", "non_AMI_CAD"):
        r = age_rows[group]
        shown = f"{float(r['median']):.1f} [{float(r['q1']):.1f}, {float(r['q3']):.1f}]"
        add("Table 1", f"Age median [IQR] {group}", "Table1_age_rebuilt_source.csv",
            f"N={r['age_available_n']}/{r['analysis_n']}; {r['median']};{r['q1']};{r['q3']}", shown)
    add("Table 1", "Age p-value", "Table1_age_rebuilt_source.csv", age_rows["overall"]["p_value"],
        f"{float(age_rows['overall']['p_value']):.3f}")

    # Verify Table 1 cell-by-cell against the locked source, except age which
    # intentionally uses the reconstructed canonical vector above.
    table1_path = next((OUT / "05_SOURCE_TRACEABILITY/source_inputs").glob("*FINAL_TABLE1_LOCK.csv"))
    table1_lock = read_csv(table1_path)
    table1_doc = Document(manuscript).tables[0]
    table1_by_label = {r.cells[0].text.strip(): r for r in table1_doc.rows[1:]}
    label_map = {"Male": "Male sex"}
    source_cols = ("overall", "definite_ami", "definite_non_ami_cad", "p_value")
    manuscript_cols = (2, 3, 4, 5)
    for lock_row in table1_lock:
        label = label_map.get(lock_row["variable"], lock_row["variable"])
        table_row = table1_by_label[label]
        if lock_row["variable"] == "Age":
            expected = (
                f"{float(age_rows['overall']['median']):.1f} [{float(age_rows['overall']['q1']):.1f}, {float(age_rows['overall']['q3']):.1f}]",
                f"{float(age_rows['AMI']['median']):.1f} [{float(age_rows['AMI']['q1']):.1f}, {float(age_rows['AMI']['q3']):.1f}]",
                f"{float(age_rows['non_AMI_CAD']['median']):.1f} [{float(age_rows['non_AMI_CAD']['q1']):.1f}, {float(age_rows['non_AMI_CAD']['q3']):.1f}]",
                f"{float(age_rows['overall']['p_value']):.3f}",
            )
            source_file = "Table1_age_rebuilt_source.csv"
        else:
            expected = tuple(
                (f"{float(lock_row[key]):.3f}" if key == "p_value" and not lock_row[key].startswith("<")
                 else lock_row[key]) for key in source_cols
            )
            source_file = table1_path.relative_to(OUT).as_posix()
        for key, col, expected_value in zip(source_cols, manuscript_cols, expected):
            cell_text = table_row.cells[col].text.strip()
            age_group = {
                "overall": "overall",
                "definite_ami": "AMI",
                "definite_non_ami_cad": "non_AMI_CAD",
                "p_value": "overall",
            }[key]
            if lock_row["variable"] == "Age" and key != "p_value":
                age_source_row = age_rows[age_group]
                source_value = (
                    f"{age_source_row['median']} [{age_source_row['q1']}, "
                    f"{age_source_row['q3']}] (N={age_source_row['age_available_n']}/"
                    f"{age_source_row['analysis_n']})"
                )
            elif lock_row["variable"] == "Age":
                source_value = age_rows["overall"]["p_value"]
            else:
                source_value = lock_row[key]
            add(f"Table 1 / {label}", key, source_file,
                source_value, expected_value, evidence_text=cell_text)

    cohort = {r["stage"]: r for r in read_csv(W["cohort"])}
    multiplicity = {int(r["rows_per_patient"]): r for r in read_csv(W["multiplicity"])}
    flow_checks = [
        ("raw_source_rows", "2,548", "Raw source records"),
        ("unique_patients_after_frozen_row_selection", "2,279", "Unique patients"),
        ("discharge_diagnosis_available", "1,944", "Discharge diagnosis available"),
        ("definite_AMI_phenotype_before_core_filter", "456", "Definite AMI before core CBC filter"),
        ("definite_non_AMI_CAD_before_core_filter", "1,372", "Definite non-AMI CAD before core CBC filter"),
        ("primary_analysis", "1,820", "Primary cohort"),
    ]
    for stage, display, label in flow_checks:
        source_row = cohort[stage]
        add("Results / cohort flow", label, W["cohort"].relative_to(ROOT).as_posix(),
            source_row["N"], display)
    for rows_per_patient, display, label in ((1, "2,010", "Patients with one source record"),
                                               (2, "269", "Patients with two source records")):
        add("Results / cohort flow", label, W["multiplicity"].relative_to(ROOT).as_posix(),
            multiplicity[rows_per_patient]["patients"], display)

    clinical_rows = [r for r in read_csv(W["clinical"])
                     if r["record_type"] == "model_performance" and
                     r["cohort_version"] == "fibrinogen_complete_case_1705"]
    for label, model in (("Clinical", "Clinical only"), ("Clinical+PIV", "Clinical + PIV"),
                         ("Clinical+Core-7", "Clinical + Core-7"),
                         ("Clinical+Enhanced", "Clinical + Enhanced")):
        r = row_by(clinical_rows, model=model)
        add("Results / clinical comparison", f"{label} AUC (95% CI)",
            W["clinical"].relative_to(ROOT).as_posix(),
            f"{r['AUC']} [{r['AUC_CI_lower']},{r['AUC_CI_upper']}]",
            ci3(r, "AUC", "AUC_CI_lower", "AUC_CI_upper"), evidence_text=workbook_text)

    for analysis in ("Symmetric +/-182-day guard-band", "High-specificity phenotype", "Conservative +/-365-day buffer"):
        for model in ("PIV", "Core-7", "Enhanced"):
            r = row_by(figure5, analysis=analysis, model=model)
            add("Table 3 / Figure 5", f"{analysis}: {model} AUC (95% CI)",
                W["figure5"].relative_to(ROOT).as_posix(), f"{r['AUC']} [{r['CI_lower']},{r['CI_upper']}]",
                ci3(r, "AUC", "CI_lower", "CI_upper"))
    write_csv(OUT / "05_SOURCE_TRACEABILITY/WP3_NUMERICAL_CROSSCHECK.csv", rows)
    return rows


def write_final_guides():
    (OUT / "06_SUBMISSION_GUIDE/BMC_REVISION_UPLOAD_GUIDE.md").write_text(
        "# BMC Revision Upload Guide\n\n"
        "This is a draft for author and reviewer review, not a portal-ready package. The exact editorial and reviewer reports and portal-integrated submitted manuscript were not recovered; compare every response with the original reports before upload.\n\n"
        "If approved after source reconciliation, upload the clean manuscript DOCX, point-by-point response DOCX, revised supplementary DOCX, and revised supplementary XLSX. Upload the marked manuscript only if the portal requests or permits it. Upload separate figures only if the portal requires them. This workflow does not submit to BMC.\n",
        encoding="utf-8")
    (OUT / "06_SUBMISSION_GUIDE/FINAL_UPLOAD_CHECKLIST.md").write_text(
        "# Final Upload Checklist\n\n"
        "- [ ] Recover the exact editor letter, Reviewer 1 report, Reviewer 2 report, and portal-integrated manuscript.\n"
        "- [ ] Reconcile every summary, response, and page location against the verbatim reports.\n"
        "- [ ] Author verifies cohort flow, age row, ethics waiver, numerical results, and limitations.\n"
        "- [ ] Confirm portal requirements for clean/marked manuscript, response, supplementary files, and separate figures.\n"
        "- [ ] Review the highlighted copy; it uses highlighting and does not fabricate tracked-change metadata.\n"
        "- [ ] Submit only after author approval. No automatic submission or PR merge was performed.\n",
        encoding="utf-8")


def finalize_package(args):
    if PACKAGE_ZIP.exists():
        raise FileExistsError(f"Refusing to overwrite archive: {PACKAGE_ZIP}")
    clean = OUT / "01_MANUSCRIPT/HITS_BMC_Revised_Manuscript_CLEAN.docx"
    marked = OUT / "01_MANUSCRIPT/HITS_BMC_Revised_Manuscript_MARKED.docx"
    supp = OUT / "03_SUPPLEMENT/Additional_file_1_Revised_Supplementary_Methods_and_Figures.docx"
    workbook = OUT / "03_SUPPLEMENT/Additional_file_2_Revised_Supplementary_Tables.xlsx"
    response = OUT / "02_RESPONSE/HITS_BMC_Point_by_Point_Response.docx"
    clean_pdf = OUT / "01_MANUSCRIPT/HITS_BMC_Revised_Manuscript_CLEAN.pdf"
    response_pdf = OUT / "02_RESPONSE/HITS_BMC_Point_by_Point_Response.pdf"
    supp_pdf = OUT / "07_QA/render_supplement/Additional_file_1_Revised_Supplementary_Methods_and_Figures.pdf"
    required = (clean, marked, supp, workbook, response, clean_pdf, response_pdf, supp_pdf)
    missing = [p.name for p in required if not p.is_file()]
    if missing:
        raise FileNotFoundError("Missing required deliverables/renders: " + ", ".join(missing))

    manuscript_text, references, cited = run_claim_reference_audit(clean)
    supplement_text = extract_docx_text(supp)
    internal_terms = ("WP1", "WP2", "canonical", "legacy", "freeze", "gate", "Codex", "GPT", "Claude", "V0.2", "V0.3")
    term_rows = [
        {"term": term,
         "manuscript_occurrences": len(re.findall(
             rf"(?<!\w){re.escape(term)}(?!\w)", manuscript_text, flags=re.IGNORECASE)),
         "supplement_occurrences": len(re.findall(
             rf"(?<!\w){re.escape(term)}(?!\w)", supplement_text, flags=re.IGNORECASE))}
        for term in internal_terms
    ]
    write_csv(OUT / "07_QA/WP3_INTERNAL_TERM_SCAN.csv", term_rows)
    ethics_terms = (
        "Ethics Committee of Xinjiang Medical University (approval no. K202602-10)",
        "requirement for informed consent was waived by the Ethics Committee",
        "All methods were performed in accordance with relevant guidelines and regulations",
    )
    ethics_pass = all(term in manuscript_text for term in ethics_terms) and not re.search(
        r"all participants provided informed consent", manuscript_text, flags=re.IGNORECASE)
    authors = ("Jibulang Ainiwaer", "Mubareke Pidamaimati", "Qian Xie", "Qixing Pi", "Wei Ji",
               "Binbin Fang", "Fen Liu", "Long Zhao", "Xiaomei Li", "Yining Yang")
    author_line = next((p.text for p in Document(clean).paragraphs if authors[0] in p.text), "")
    positions = [author_line.find(name) for name in authors]
    author_pass = bool(author_line) and all(p >= 0 for p in positions) and positions == sorted(positions)
    all_text = manuscript_text + "\n" + supplement_text
    funding_pass = all(x in all_text for x in ("82560073", "2023TSYCLJ0035", "2023D01D12"))
    data_pass = "participant-privacy restrictions" in manuscript_text and "reasonable request" in manuscript_text
    ai_pass = "English-language polishing" in manuscript_text and "code drafting" in manuscript_text
    missing_citations = sorted(cited - set(references))
    orphan_references = sorted(set(references) - cited)
    figure_pass = all(f"Figure {i}." in manuscript_text for i in range(1, 6))
    table_pass = all(f"Table {i}." in manuscript_text for i in range(1, 4))
    internal_pass = all(r["manuscript_occurrences"] == 0 and r["supplement_occurrences"] == 0 for r in term_rows)
    crosschecks = build_numerical_crosscheck(clean, supp)
    numerical_pass = all(r["match"] == "YES" for r in crosschecks)
    gate = "WP3_HOLD_COMMENT_UNANSWERED"

    output_figures = []
    for stem in ("Figure1_Revised", "Figure2_Revised", "Figure3_Revised", "Figure4_Revised", "Figure5_Revised",
                 "FigureS1_DCA", "FigureS2_HighSpecificityFlow", "FigureS3_Core7Correlation", "FigureS4_Domains"):
        for ext in ("pdf", "tiff", "png"):
            path = OUT / "04_FIGURES" / f"{stem}.{ext}"
            if not path.is_file() or path.stat().st_size == 0:
                raise FileNotFoundError(f"Missing figure export: {path.name}")
            output_figures.append(path)
    if not numerical_pass or not ethics_pass or not author_pass or not funding_pass or not data_pass or not ai_pass:
        raise RuntimeError("A required final QA check failed; see 07_QA reports.")
    if missing_citations or orphan_references or not figure_pass or not table_pass or not internal_pass:
        raise RuntimeError("Citation, caption, or internal-term QA failed; see 07_QA reports.")

    qa_rows = [
        ("Numerical cross-check", "PASS", f"{len(crosschecks)} registered values matched"),
        ("Ethics statement and consent correction", "PASS", "exact approval/waiver/guidelines text present; universal-consent statement absent"),
        ("Author order", "PASS", "all ten authors retained in order"),
        ("Funding", "PASS", "three locked grant numbers present"),
        ("Data availability", "PASS", "restricted-data and reasonable-request language retained"),
        ("AI disclosure", "PASS", "existing disclosure boundary retained"),
        ("Citations/references", "PASS", "no missing citations or orphan references"),
        ("Main figure/table captions", "PASS", "Figures 1-5 and Tables 1-3 present"),
        ("Internal terms", "PASS", "blocked project/version terms absent from manuscript and supplementary text"),
        ("Raw patient-level privacy", "PASS", "restricted source excluded; copied analysis data are aggregate"),
        ("Exact editor/reviewer sources", "HOLD", "verbatim editor and reviewer reports not recovered"),
        ("Portal manuscript snapshot", "HOLD", "exact portal-integrated manuscript not recovered"),
        ("Response-to-source alignment", "HOLD", "responses use explicitly marked non-verbatim summaries"),
        ("Final Gate", gate, "compare with original journal source comments before submission"),
    ]
    write_csv(OUT / "07_QA/WP3_QA_CHECKS.csv",
              [{"check": a, "status": b, "evidence": c} for a, b, c in qa_rows])
    write_final_guides()
    qa_report = f"""# WP3 Final QA Report

**Technical package QA:** PASS for registered numerical checks, declarations, citations, table/figure captions, and internal-term scan.

**Final Gate:** {gate}

The exact editor letter, Reviewer 1 and Reviewer 2 reports, and portal-integrated submitted manuscript were not recovered. The response uses non-verbatim summaries, so alignment to the journal's exact requests cannot be certified. This is a source blocker, not a formatting issue. Do not submit before recovering and reconciling the original sources.

## Comment Coverage

- Editor: 10/10 summary items addressed with explicit source limitation.
- Reviewer 1: 10/10 summary items addressed with explicit source limitation.
- Reviewer 2: 1/1 summary item addressed with explicit source limitation.

These counts do not mean verified closure against the verbatim reports.

## Scientific Boundaries

- Single-center retrospective CAD cohort; AMI phenotype discrimination only.
- No future-event, MACE, prospective, external, formal temporal-validation, diagnostic-readiness, triage, or deployment claim.
- CBC timing relative to the index angiography hospitalization is unverified for all 1,820 selected patients.
- AMI phenotype uses discharge-diagnosis text without independent uniform adjudication.
- Clinical baseline: age, sex, hypertension, and diabetes; smoking was unavailable.
- Decision-curve analysis is exploratory; no harmonized troponin comparator was fitted.
- Restricted source rows, patient identifiers, and patient-level predictions are excluded.

## Supporting QA

See WP3_QA_CHECKS.csv, WP3_NUMERICAL_CROSSCHECK.csv, WP3_CLAIM_AUDIT.csv, WP3_CITATION_REFERENCE_AUDIT.csv, WP3_INTERNAL_TERM_SCAN.csv, and WP3_MANUSCRIPT_PAGE_MAP.csv.

The marked manuscript uses highlighting, not fabricated Word tracked-change metadata. No BMC submission or PR merge was performed.
"""
    (OUT / "07_QA/WP3_FINAL_QA_REPORT.md").write_text(qa_report, encoding="utf-8")

    sanitize_public_paths()
    checksum_rows = []
    for path in sorted(OUT.rglob("*")):
        if path.is_file() and path.name != "WP3_PACKAGE_SHA256SUMS.csv":
            checksum_rows.append({"path": path.relative_to(OUT).as_posix(),
                                  "SHA256": sha256(path), "bytes": path.stat().st_size})
    write_csv(OUT / "07_QA/WP3_PACKAGE_SHA256SUMS.csv", checksum_rows)
    with zipfile.ZipFile(PACKAGE_ZIP, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
        for path in sorted(OUT.rglob("*")):
            if path.is_file():
                zf.write(path, Path(OUT.name) / path.relative_to(OUT))
    with zipfile.ZipFile(PACKAGE_ZIP) as zf:
        bad_member = zf.testzip()
        if bad_member:
            raise RuntimeError(f"ZIP integrity failure: {bad_member}")
        names = zf.namelist()
        if any(re.search(r"patient[_-]?id|research_patient_id|patient_sn|住院号", n, flags=re.IGNORECASE) for n in names):
            raise RuntimeError("Potential patient-level file name detected in ZIP.")
    digest = sha256(PACKAGE_ZIP)
    (ROOT / "HITS_BMC_MINOR_REVISION_WP3_FINAL.zip.sha256").write_text(
        f"{digest}  {PACKAGE_ZIP.name}\n", encoding="ascii")
    append_run_log(f"FINALIZE_PASS gate={gate}; crosschecks={len(crosschecks)}; zip_sha256={digest}")
    print(json.dumps({"phase": "finalize", "gate": gate, "zip": str(PACKAGE_ZIP),
                      "sha256": digest, "numeric_pass": numerical_pass, "crosschecks": len(crosschecks)}))


def main():
    parser = argparse.ArgumentParser(description="Build the bounded HITS BMC WP3 revision package.")
    parser.add_argument("--phase", required=True, choices=("prepare", "assemble", "response", "finalize"))
    parser.add_argument("--manuscript")
    parser.add_argument("--supplement")
    parser.add_argument("--workbook")
    parser.add_argument("--table1-source")
    parser.add_argument("--restricted-master")
    parser.add_argument("--restricted-master-sha256", default=EXPECTED_MASTER_SHA256)
    parser.add_argument("--manuscript-pdf")
    parser.add_argument("--supplement-pdf")
    args = parser.parse_args()
    if args.phase == "prepare":
        if not all((args.manuscript, args.supplement, args.workbook, args.table1_source, args.restricted_master)):
            parser.error("prepare requires manuscript, supplement, workbook, table1-source, and restricted-master paths")
        prepare_package(args)
    elif args.phase == "assemble":
        if not args.manuscript or not args.supplement:
            parser.error("assemble requires --manuscript and --supplement")
        assemble_package(args)
    elif args.phase == "response":
        if not args.manuscript_pdf or not args.supplement_pdf:
            parser.error("response requires --manuscript-pdf and --supplement-pdf")
        response_package(args)
    elif args.phase == "finalize":
        finalize_package(args)


if __name__ == "__main__":
    main()
