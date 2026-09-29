"""Read-only WP1R-A audit of the four named source files and source linkage."""
import csv
import hashlib
import json
import os
import re
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

PROJECT=Path(os.environ.get('HITS_PROJECT_ROOT',Path(__file__).resolve().parents[1]))
OUT=PROJECT/'10_ORIGINAL_SOURCE_RECOVERY'
DOC=Path(os.environ.get('HITS_DOCUMENTS_DIR',Path.home()/'Documents'))
DOWN=Path(os.environ.get('HITS_DOWNLOADS_DIR',Path.home()/'Downloads'))
DS=Path(os.environ.get('HITS_SOURCE_RAW_DIR',DOC/'颈动脉粥样斑块/Dual_Score_V0.3/source_raw'))
CLEAN=Path(os.environ.get('HITS_MASTER_CSV',DOC/'颈动脉粥样斑块MACE事件/下载归档/2026-08-16_Article2_投稿预审与队列溯源/TwoPapers_Article2_R代码与图表/TwoPapers_R_Work_V0.3/data_raw/ehr_extracted/master_cohort_cleaned.csv'))

FILES={
 'primary_mother':'冠心病+造影手术+20-26（减少附属）.csv',
 'adverse_plus':'冠心病+造影手术+20-26+不良（减少附属）.csv',
 'adverse_dash':'冠心病+造影手术+20-26-不良（减少附属）.csv',
 'adverse_xlsx':'冠心病+造影手术+20-26年+不良.xlsx'}
RENAMES={'primary_mother':['冠心病+造影手术+20-26（减少附属）(1).csv'],
         'adverse_plus':['冠心病+造影手术+20-26+不良（减少附属）(1).csv'],
         'adverse_dash':['冠心病+造影手术+20-26-不良（减少附属）(1).csv'],
         'adverse_xlsx':['冠心病+造影手术+20-26年+不良(1).xlsx']}


def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()


def csvout(name,rows):
    p=OUT/name;p.parent.mkdir(parents=True,exist_ok=True)
    (rows if isinstance(rows,pd.DataFrame) else pd.DataFrame(rows)).to_csv(p,index=False,lineterminator='\n')


def locate(key):
    names={FILES[key],*RENAMES[key]}
    found=[]
    for base in (DOWN,DS):
        if base.is_dir():
            for name in names:
                found.extend(p for p in base.rglob(name) if p.is_file())
    return list(dict.fromkeys(found))


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    rows=[];found={}
    for key,official in FILES.items():
        paths=locate(key);found[key]=paths
        if not paths:
            rows.append(dict(file_role=key,official_filename=official,copy_location='NOT_FOUND',exists='NO',sha256='',bytes='',rows='',columns='',unique_patient_sn='',three_header_rows='UNVERIFIED',status='MISSING'))
            continue
        hashes=set()
        for p in paths:
            h=sha(p);hashes.add(h);nrow=ncol=nids='';header='NOT_APPLICABLE'
            if p.suffix.lower()=='.csv':
                with p.open(encoding='utf-8-sig',newline='') as f:
                    reader=csv.reader(f);hs=[next(reader,None) for _ in range(3)];nrow=sum(1 for _ in reader)
                ncol=len(hs[0] or []);header='YES' if len({len(x or []) for x in hs})==1 and ncol>0 else 'NO'
                data=pd.read_csv(p,header=2,dtype='string',keep_default_na=False,low_memory=False)
                if 'patient_sn' in data:nids=data.patient_sn.nunique()
            else:
                book=load_workbook(p,read_only=True,data_only=True)
                for sh in book:
                    it=sh.iter_rows(values_only=True);hs=[next(it,None) for _ in range(3)];count=sum(1 for _ in it)
                    location='Downloads' if DOWN in p.parents else 'Documents/source_raw'
                    rows.append(dict(file_role=key,official_filename=official,copy_location=location,sheet=sh.title,exists='YES',sha256=h,bytes=p.stat().st_size,rows=count,columns=sh.max_column,unique_patient_sn='',three_header_rows='YES' if all(x is not None for x in hs) else 'NO',status='FOUND_COPY'))
                continue
            location='Downloads' if DOWN in p.parents else 'Documents/source_raw'
            rows.append(dict(file_role=key,official_filename=official,copy_location=location,exists='YES',sha256=h,bytes=p.stat().st_size,rows=nrow,columns=ncol,unique_patient_sn=nids,three_header_rows=header,status='FOUND_COPY'))
        if len(hashes)>1:raise AssertionError(f'Different copies found for {key}')
    csvout('ORIGINAL_FOUR_FILE_INVENTORY.csv',rows)

    mother=next((p for p in found['primary_mother'] if p.suffix.lower()=='.csv'),None)
    if mother is None:raise FileNotFoundError('Primary mother CSV not found')
    raw=pd.read_csv(mother,header=2,dtype='string',keep_default_na=False,low_memory=False)
    if raw.shape!=(2548,431) or raw.patient_sn.nunique()!=2279:raise AssertionError('Primary mother structure/count does not match expected anchor')
    clean=pd.read_csv(CLEAN,dtype='string',keep_default_na=False,low_memory=False)
    if set(raw.patient_sn)!=set(clean.patient_sn):raise AssertionError('Primary and cleaned patient-key sets differ')
    clean=clean.drop_duplicates('patient_sn').set_index('patient_sn')

    mapped={
      'discharge_diagnosis':('discharge_diagnosis','discharge_diagnosis'),
      'admission_date':('admission_date','admission_date'),
      'wbc_value':('lab_blood_routine_examination_WBC#_test_result','lab_blood_routine_examination_WBC_test_result'),
      'neut_value':('lab_blood_routine_examination_Neut#_test_result','lab_blood_routine_examination_Neut_test_result__dup02'),
      'lymph_value':('lab_blood_routine_examination_Lymph#_test_result','lab_blood_routine_examination_Lymph_test_result__dup02'),
      'mono_value':('lab_blood_routine_examination_Mono#_test_result','lab_blood_routine_examination_Mono_test_result__dup02'),
      'plt_value':('lab_blood_routine_examination_PLT#_test_result','lab_blood_routine_examination_PLT_test_result'),
      'mpv_value':('lab_blood_routine_examination_MPV_test_result','lab_blood_routine_examination_MPV_test_result'),
      'rdw_value':('lab_blood_routine_examination_RDW-CV_test_result','lab_blood_routine_examination_RDW_CV_test_result'),
      'hb_value':('lab_blood_routine_examination_Hb_test_result','lab_blood_routine_examination_Hb_test_result'),
      'fbg_value':('lab_cruor_testing_Fbg_test_result','lab_cruor_testing_Fbg_test_result'),
      'wbc_time':('lab_blood_routine_examination_WBC#_test_time','lab_blood_routine_examination_WBC_test_time'),
      'neut_time':('lab_blood_routine_examination_Neut#_test_time','lab_blood_routine_examination_Neut_test_time__dup02'),
      'lymph_time':('lab_blood_routine_examination_Lymph#_test_time','lab_blood_routine_examination_Lymph_test_time__dup02'),
      'mono_time':('lab_blood_routine_examination_Mono#_test_time','lab_blood_routine_examination_Mono_test_time__dup02'),
      'plt_time':('lab_blood_routine_examination_PLT#_test_time','lab_blood_routine_examination_PLT_test_time'),
      'mpv_time':('lab_blood_routine_examination_MPV_test_time','lab_blood_routine_examination_MPV_test_time'),
      'rdw_time':('lab_blood_routine_examination_RDW-CV_test_time','lab_blood_routine_examination_RDW_CV_test_time'),
      'hb_time':('lab_blood_routine_examination_Hb_test_time','lab_blood_routine_examination_Hb_test_time'),
      'fbg_time':('lab_cruor_testing_Fbg_test_time','lab_cruor_testing_Fbg_test_time')}
    names=set(raw.columns)
    if not all(a in names and b in clean.columns for a,b in mapped.values()):raise AssertionError('A mapped source/clean field is missing')
    id_rows=raw.groupby('patient_sn',sort=False).size();repeat=raw[raw.patient_sn.isin(id_rows[id_rows.eq(2)].index)]
    pairs=[]
    for field in ['基线_inpatient_no_id','inpatient_no_id','admission_date']:
        counts=repeat.groupby('patient_sn')[field].apply(lambda z:z.replace('',pd.NA).dropna().nunique())
        nonempty=repeat.groupby('patient_sn')[field].apply(lambda z:int(z.ne('').sum()))
        pairs.append(dict(field=field,repeated_patient_pairs=269,pairs_both_rows_nonmissing=int(nonempty.eq(2).sum()),pairs_one_row_nonmissing=int(nonempty.eq(1).sum()),pairs_distinct_values=int(counts.gt(1).sum()),pairs_same_single_value=int(counts.eq(1).sum()),pairs_no_value=int(counts.eq(0).sum())))
    csvout('REPEATED_PATIENT_TOKEN_AUDIT.csv',pairs)
    # Check values at patient level only; never persist a patient key or row.
    cross=[]
    ra=raw[['patient_sn']+[v[0] for v in mapped.values()]].rename(columns={v[0]:k+'_raw' for k,v in mapped.items()})
    cb=clean[[v[1] for v in mapped.values()]].rename(columns={v[1]:k+'_clean' for k,v in mapped.items()})
    joined=ra.merge(cb,left_on='patient_sn',right_index=True,how='left',validate='many_to_one')
    for group,keys in [('diagnosis',['discharge_diagnosis']),('admission_date',['admission_date']),('core_CBC',['neut_value','lymph_value','mono_value','plt_value','mpv_value','rdw_value','hb_value']),('CBC_plus_WBC',['wbc_value','neut_value','lymph_value','mono_value','plt_value','mpv_value','rdw_value','hb_value']),('CBC_timestamp_set',['wbc_time','neut_time','lymph_time','mono_time','plt_time','mpv_time','rdw_time','hb_time']),('fibrinogen_value_time',['fbg_value','fbg_time']),('all_mapped_fields',list(mapped))]:
        eq=pd.Series(True,index=joined.index)
        for k in keys:eq &= joined[k+'_raw'].eq(joined[k+'_clean'])
        n=int(joined.assign(_eq=eq).groupby('patient_sn')._eq.any().sum())
        cross.append(dict(field_set=group,patients_with_exact_source_row_match=n,patients_total=2279,comparison='Any source row for same patient; literal values, no patient key exported'))
    csvout('RAW_TO_ANALYSIS_SOURCE_CROSSWALK.csv',cross)

    # All candidate dates are reduced to counts and years before publication.
    date_rows=[]; field_rows=[]
    headers=[list(map(str,x)) for x in zip(*[["" for _ in range(431)]]*3)] if False else None
    with mother.open(encoding='utf-8-sig',newline='') as f:
        reader=csv.reader(f);hdr=[next(reader) for _ in range(3)]
    for i in range(431):
        module,label,name=(hdr[j][i] if i<len(hdr[j]) else '' for j in range(3))
        text=' '.join([module,label,name])
        field_rows.append(dict(column_index=i+1,module=module,chinese_label=label,machine_name=name))
        if re.search(r'date|time|日期|时间|test_time|检验日期',text,re.I):
            s=pd.to_datetime(raw.iloc[:,i].replace('',pd.NA),errors='coerce',format='mixed')
            date_rows.append(dict(column_index=i+1,module=module,chinese_label=label,machine_name=name,nonmissing_n=int(raw.iloc[:,i].ne('').sum()),parseable_n=int(s.notna().sum()),earliest_year=int(s.dt.year.min()) if s.notna().any() else '',latest_year=int(s.dt.year.max()) if s.notna().any() else '',before_2019_n=int(s.dt.year.lt(2019).sum())))
    csvout('ORIGINAL_MOTHER_FIELD_INVENTORY.csv',field_rows)
    csvout('ORIGINAL_MOTHER_DATE_FIELD_SUMMARY.csv',date_rows)
    wbc=raw['lab_blood_routine_examination_WBC#_test_time'].replace('',pd.NA)
    wp=pd.to_datetime(wbc,errors='coerce',format='mixed')
    fbg=pd.to_datetime(raw['lab_cruor_testing_Fbg_test_time'].replace('',pd.NA),errors='coerce',format='mixed')
    admission=pd.to_datetime(raw.admission_date.replace('',pd.NA),errors='coerce',format='mixed')
    timing=[]
    for label,lab in [('WBC_time',wp),('Fbg_time',fbg)]:
        ok=lab.notna()&admission.notna();delta=(lab[ok]-admission[ok]).dt.total_seconds()/86400
        timing.append(dict(comparison=f'{label}_minus_admission',available_n=int(ok.sum()),median_days=float(delta.median()) if len(delta) else '',before_event_n=int(delta.lt(0).sum()),after_event_n=int(delta.gt(0).sum()),same_calendar_day_n=int((lab[ok].dt.date==admission[ok].dt.date).sum()),semantic='Descriptive dates only; same-encounter relationship is not evidenced'))
    csvout('ORIGINAL_DATE_RELATIONSHIPS.csv',timing)
    summary=dict(primary_raw_rows=len(raw),columns=raw.shape[1],unique_patient_sn=int(raw.patient_sn.nunique()),patient_set_exact_match_clean_master=True,clean_master_rows=len(pd.read_csv(CLEAN,usecols=['patient_sn'],dtype='string')),clean_master_unique_patients=int(clean.index.nunique()),repeated_patients=int(id_rows.gt(1).sum()),surplus_rows=int((id_rows-1).sum()),one_row_patients=int(id_rows.eq(1).sum()),two_row_patients=int(id_rows.eq(2).sum()),patients_ge3=int(id_rows.gt(2).sum()),repeated_groups_same_baseline_admission_token=int(next(x['pairs_same_single_value'] for x in pairs if x['field']=='基线_inpatient_no_id')),repeated_groups_multiple_current_admission_numbers=int(next(x['pairs_distinct_values'] for x in pairs if x['field']=='inpatient_no_id')),wbc_test_time_n=int(wp.notna().sum()),wbc_before_2019_n=int(wp.dt.year.lt(2019).sum()),fbg_test_time_n=int(fbg.notna().sum()),fbg_before_2019_n=int(fbg.dt.year.lt(2019).sum()),discharge_date_field_present=any(re.search(r'discharge_date|discharge_time|出院日期|出院时间',c,re.I) for c in raw.columns),cag_pci_procedure_date_field_present=any(re.search(r'angiograph|\bcag\b|\bpci\b|造影|procedure_date|手术日期',c,re.I) for c in raw.columns),models_run=False)
    (OUT/'SOURCE_AUDIT_METRICS.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(summary,ensure_ascii=False))


if __name__=='__main__':main()
